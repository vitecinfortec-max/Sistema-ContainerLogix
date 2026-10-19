"""Relatório de Compra e Venda de Container: estoque (comprados sem venda,
dias parado), vendas com margem sobre a compra vinculada, a receber, totais
por representante, pendências de cadastro, filtros e arquivos."""
import io
from datetime import datetime, timedelta

from openpyxl import load_workbook

from reports import BRT_TZ
from test_gestao_container import _purchase, _sale


def _set(api, collection, doc_id, **fields):
    api.run(api.db[collection].update_one({"id": doc_id}, {"$set": fields}))


def test_stock_sales_margin_and_pending(api):
    old = _purchase(api, "ABCU1111111", purchase_value=5000, sale_value=8000)
    _set(api, "container_purchases", old["id"], entry_date=(datetime.now(BRT_TZ) - timedelta(days=120)).isoformat())
    _purchase(api, "ABCU2222222", purchase_value=6000)
    sold = _sale(api, "ABCU2222222", sale_value=9000)                    # vincula a compra de 6000
    _sale(api, "ABCU3333333", sale_value=7000)                           # venda sem compra
    _sale(api, "ABCU4444444")                                            # sem compra e sem valor
    assert api.put(f"/api/container-sales/{sold['id']}/status", params={"status": "RECEBIDO"}).status_code == 200

    data = api.get("/api/reports/containers").json()
    s = data["summary"]
    assert s["stock_count"] == 1 and s["stock_value"] == 5000 and s["stock_over_90"] == 1
    assert data["stock"][0]["container_number"] == "ABCU1111111" and data["stock"][0]["days_in_stock"] == 120
    assert s["sales_count"] == 3 and s["sales_total"] == 16000
    assert s["margin_total"] == 3000 and s["cost_total"] == 6000
    assert s["open_total"] == 7000  # a recebida não conta; a sem valor também não
    assert s["without_purchase"] == 2 and s["without_value"] == 1
    by_container = {x["container_number"]: x for x in data["sales"]}
    assert by_container["ABCU2222222"]["margin_pct"] == 50.0
    assert by_container["ABCU3333333"]["margin"] is None


def test_period_and_representative_filters(api):
    a = _sale(api, "ABCU1111111", sale_value=1000)
    _sale(api, "ABCU2222222", sale_value=2000)
    _set(api, "container_sales", a["id"], created_at="2026-01-10T15:00:00+00:00", representative_name="João Rep")
    jan = api.get("/api/reports/containers", params={"date_from": "2026-01-01", "date_to": "2026-01-31"}).json()
    assert [x["container_number"] for x in jan["sales"]] == ["ABCU1111111"]
    assert api.get("/api/reports/containers", params={"representative": "joão"}).json()["summary"]["sales_total"] == 1000
    reps = api.get("/api/reports/containers").json()["representatives"]
    assert {r["representative_name"] for r in reps} == {"João Rep", "(sem representante)"}


def test_files_and_admin_only(api):
    _purchase(api, "ABCU1111111", purchase_value=5000)
    _sale(api, "ABCU2222222", sale_value=7000)
    assert api.get("/api/reports/containers/pdf").content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(api.get("/api/reports/containers/excel").content))
    assert wb.sheetnames == ["Vendas", "Estoque", "Por representante", "Pendências"]
    values = [c for row in wb["Estoque"].iter_rows(values_only=True) for c in row if c is not None]
    assert "ABCU1111111" in values
    assert api.get("/api/reports/containers", as_user="operator").status_code == 403
