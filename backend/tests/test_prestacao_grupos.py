"""Prestação de Contas: grupo de custo em cada compra, cadastro dos grupos,
dashboard de custo por grupo e o relatório com os maiores custos."""
import io

from openpyxl import load_workbook


def _groups(api):
    r = api.get("/api/expense-groups")
    assert r.status_code == 200, r.text
    return {g["name"]: g for g in r.json()}


def _purchase(group_id, date, amount, supplier="Posto Central", observation=None):
    return {"supplier_name": supplier, "purchase_date": date, "amount": amount, "group_id": group_id, "observation": observation}


def _create(api, purchases, deposits=None, start="2026-10-01", end="2026-10-31"):
    r = api.post("/api/expense-reports", json={"period_start": start, "period_end": end, "purchases": purchases,
                                               "deposits": deposits or []})
    assert r.status_code == 200, r.text
    return r.json()


def test_default_groups_and_crud(api):
    groups = _groups(api)
    assert "Combustível" in groups and "Alimentação" in groups and len(groups) == 10
    assert len(_groups(api)) == 10  # não duplica ao listar de novo
    new = api.post("/api/expense-groups", json={"name": "Lavagem de veículos"}).json()
    assert api.post("/api/expense-groups", json={"name": "lavagem de VEÍCULOS"}).status_code == 400
    assert api.post("/api/expense-groups", json={"name": "  "}).status_code == 400
    rep = _create(api, [_purchase(new["id"], "2026-10-05", 80)])
    assert rep["purchases"][0]["group_name"] == "Lavagem de veículos"
    # renomear acompanha as compras já lançadas; usado não exclui, só inativa
    api.put(f"/api/expense-groups/{new['id']}", json={"name": "Lavagem", "status": "ATIVO"})
    assert api.get(f"/api/expense-reports/{rep['id']}").json()["purchases"][0]["group_name"] == "Lavagem"
    r = api.delete(f"/api/expense-groups/{new['id']}")
    assert r.status_code == 400 and "Inativo" in r.json()["detail"]
    assert api.put(f"/api/expense-groups/{new['id']}", json={"name": "Lavagem", "status": "INATIVO"}).json()["status"] == "INATIVO"
    unused = api.post("/api/expense-groups", json={"name": "Temporário"}).json()
    assert api.delete(f"/api/expense-groups/{unused['id']}").status_code == 200
    # grupo inexistente é recusado
    r = api.post("/api/expense-reports", json={"period_start": "2026-10-01", "period_end": "2026-10-31",
                                               "purchases": [_purchase("nao-existe", "2026-10-05", 10)], "deposits": []})
    assert r.status_code == 400


def test_dashboard_by_group_and_previous_period(api):
    g = _groups(api)
    fuel, food = g["Combustível"]["id"], g["Alimentação"]["id"]
    _create(api, [_purchase(fuel, "2026-09-10", 400), _purchase(food, "2026-09-12", 100)], start="2026-09-01", end="2026-09-30")
    _create(api, [_purchase(fuel, "2026-10-03", 600), _purchase(fuel, "2026-10-20", 300, supplier="Posto Rota"),
                  _purchase(food, "2026-10-05", 80), _purchase(None, "2026-10-06", 20, supplier="Papelaria")],
            deposits=[{"amount": 500, "date": "2026-10-01", "sent_by": "Financeiro"}])
    r = api.get("/api/expense-reports/dashboard", params={"date_from": "2026-10-01", "date_to": "2026-10-31"})
    assert r.status_code == 200, r.text
    d = r.json()
    groups = {x["group_name"]: x for x in d["groups"]}
    assert d["summary"]["total_purchases"] == 1000 and d["summary"]["without_group_total"] == 20
    assert groups["Combustível"]["total"] == 900 and groups["Combustível"]["share"] == 90 and groups["Combustível"]["count"] == 2
    assert groups["Combustível"]["largest"] == 600 and groups["Combustível"]["average"] == 450
    assert d["groups"][0]["group_name"] == "Combustível"  # maior custo primeiro
    # outubro x setembro (período anterior de mesmo tamanho)
    assert groups["Combustível"]["previous_total"] == 400 and groups["Combustível"]["variation"] == 125
    assert groups["Alimentação"]["variation"] == -20 and groups["(sem grupo)"]["variation"] is None
    assert d["previous_period"]["date_from"] == "2026-08-31" and d["previous_period"]["total"] == 500
    assert d["top_purchases"][0]["amount"] == 600 and d["top_suppliers"][0]["supplier_name"] == "Posto Central"
    assert d["by_month"] == [{"month": "2026-10", "total": 1000, "Combustível": 900, "Alimentação": 80, "(sem grupo)": 20}]


def test_report_filters_by_group_and_lists_largest_costs(api):
    g = _groups(api)
    fuel, food = g["Combustível"]["id"], g["Alimentação"]["id"]
    _create(api, [_purchase(fuel, "2026-10-03", 600), _purchase(food, "2026-10-05", 80), _purchase(None, "2026-10-06", 20)],
            deposits=[{"amount": 500, "date": "2026-10-01", "sent_by": "Financeiro"}])
    only_food = api.get("/api/reports/expenses", params={"group": food}).json()
    assert only_food["summary"]["total_purchases"] == 80 and only_food["deposits"] == []
    no_group = api.get("/api/reports/expenses", params={"group": "SEM_GRUPO"}).json()
    assert no_group["summary"]["total_purchases"] == 20
    everything = api.get("/api/reports/expenses", params={"group": "TODOS"}).json()
    assert everything["summary"]["total_purchases"] == 700 and everything["top_purchases"][0]["amount"] == 600
    pdf = api.get("/api/reports/expenses/pdf", params={"group": food})
    assert pdf.content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(api.get("/api/reports/expenses/excel").content))
    sheet = wb["Por grupo"]
    names = {v for row in sheet.iter_rows(values_only=True) for v in row if isinstance(v, str)}  # a tabela começa na coluna B
    assert "Combustível" in names and "Alimentação" in names


def _pdf_text(content: bytes) -> str:
    from io import BytesIO, StringIO
    from pdfminer.high_level import extract_text_to_fp
    out = StringIO()
    extract_text_to_fp(BytesIO(content), out)
    return "".join(out.getvalue().split())


def test_pdfs_bring_the_cost_dashboard(api):
    g = _groups(api)
    rep = _create(api, [_purchase(g["Combustível"]["id"], "2026-09-03", 600), _purchase(g["Alimentação"]["id"], "2026-10-05", 80)])
    api.put(f"/api/expense-reports/{rep['id']}/status", params={"status": "CONCLUIDA"})
    one = _pdf_text(api.get(f"/api/expense-reports/{rep['id']}/pdf").content)
    assert "Custoporgrupo" in one and "Combustível" in one and "88,2%" in one
    period = _pdf_text(api.get("/api/reports/expenses/pdf").content)
    assert "Dashboard-custoporgrupo" in period and "Dashboard-totalgastopormês" in period
