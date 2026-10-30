"""Nos relatórios o motorista sai só com o primeiro e o segundo nome; os
documentos com assinatura/CPF continuam com o nome completo."""
import io

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from reports import generate_excel_report, generate_pdf_report, short_person_name

FULL_NAME = "JOSE DA SILVA SANTOS NETO"


@pytest.mark.parametrize("name, expected", [
    (FULL_NAME, "JOSE SILVA"),
    ("Maria de Fatima Souza", "Maria Fatima"),
    ("  JOAO   PEDRO   ALVES ", "JOAO PEDRO"),
    ("JOSE", "JOSE"),
    ("(sem motorista)", "(sem motorista)"),
    ("", "-"),
    (None, "-"),
])
def test_short_person_name(name, expected):
    assert short_person_name(name) == expected


def _movement():
    return {
        "transaction_id": 1, "created_at": "2026-10-30T12:00:00+00:00", "operation_type": "ENTRADA",
        "container_number": "MSCU1234565", "driver_name": FULL_NAME, "status": "VAZIO",
    }


def test_movement_pdf_shows_short_driver_name():
    text = PdfReader(io.BytesIO(generate_pdf_report([_movement()], show_container_value=False))).pages[0].extract_text()
    assert "JOSE SILVA" in text and "SANTOS" not in text


def test_movement_excel_shows_short_driver_name():
    sheet = load_workbook(io.BytesIO(generate_excel_report([_movement()]))).active
    values = [cell for row in sheet.iter_rows(values_only=True) for cell in row]
    assert "JOSE SILVA" in values and FULL_NAME not in values
