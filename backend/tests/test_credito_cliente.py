"""Crédito de Cliente: lançar crédito, dar baixa em faturas (parcial e total),
estornos, devolução, efeitos na fatura (cancelar, excluir, total menor) e o
extrato."""
import io
from datetime import timedelta

from openpyxl import load_workbook

from routers.client_credits import _today

TODAY = _today()
DAY = lambda n=0: (TODAY - timedelta(days=n)).isoformat()


def _seed(api):
    api.run(api.db.clients.insert_many([
        {"id": "cA", "name": "CLIENTE A", "cnpj": "11.111.111/0001-11"},
        {"id": "cB", "name": "CLIENTE B", "cnpj": None},
    ]))
    created = f"{DAY(10)}T15:00:00+00:00"
    api.run(api.db.movements.insert_many([
        {"id": f"m{i}", "service_value": v, "billed": True, "invoice_id": inv}
        for i, (v, inv) in enumerate([(300, "i1"), (200, "i1"), (500, "i2"), (400, "i3")])
    ]))
    api.run(api.db.invoices.insert_many([
        {"id": "i1", "invoice_number": 1, "client_name": "CLIENTE A", "client_cnpj": "11.111.111/0001-11", "movement_ids": ["m0", "m1"],
         "total_value": 500, "status": "PENDENTE", "created_at": created, "created_by": "adm", "user_name": "Admin"},
        {"id": "i2", "invoice_number": 2, "client_name": "cliente a", "client_cnpj": "11111111000111", "movement_ids": ["m2"],
         "total_value": 500, "status": "PENDENTE", "created_at": created, "created_by": "adm", "user_name": "Admin"},
        {"id": "i3", "invoice_number": 3, "client_name": "CLIENTE B", "movement_ids": ["m3"],
         "total_value": 400, "status": "PENDENTE", "created_at": created, "created_by": "adm", "user_name": "Admin"},
    ]))


def _credit(api, amount, client="cA", entry_type="CREDITO", day=None, **extra):
    return api.post("/api/client-credits/entries", json={"client_id": client, "entry_type": entry_type, "amount": amount,
                                                         "entry_date": day or DAY(), **extra})


def _apply(api, *allocations, client="cA", day=None):
    return api.post("/api/client-credits/apply", json={"client_id": client, "entry_date": day or DAY(),
                                                       "allocations": [{"invoice_id": i, "amount": a} for i, a in allocations]})


def _balance(api, client="cA"):
    return next(r["balance"] for r in api.get("/api/client-credits/accounts").json() if r["client_id"] == client)


def _invoice(api, id):
    return api.get(f"/api/invoices/{id}").json()


def test_credit_and_partial_then_full_use(api):
    _seed(api)
    assert _credit(api, 700, payment_method="PIX", reference="E123").status_code == 200
    assert _apply(api, ("i1", 200)).status_code == 200
    inv = _invoice(api, "i1")
    assert (inv["status"], inv["credit_applied"], inv["open_value"]) == ("PENDENTE", 200, 300)
    # resto da i1 + parte da i2 numa baixa só; i1 fica quitada pelo crédito
    assert _apply(api, ("i1", 300), ("i2", 150)).status_code == 200
    inv = _invoice(api, "i1")
    assert (inv["status"], inv["paid_date"], inv["paid_by_credit"]) == ("PAGO", DAY(), True)
    assert _invoice(api, "i2")["open_value"] == 350
    assert _balance(api) == 50
    row = next(r for r in api.get("/api/client-credits/accounts").json() if r["client_id"] == "cA")
    assert (row["total_credited"], row["total_used"], row["open_invoices"]) == (700, 650, 350)
    open_list = api.get("/api/client-credits/open-invoices", params={"client_id": "cA"}).json()
    assert [i["invoice_number"] for i in open_list["invoices"]] == [2] and open_list["balance"] == 50


def test_limits(api):
    _seed(api)
    _credit(api, 100)
    assert "insuficiente" in _apply(api, ("i1", 150)).json()["detail"]
    assert "outro cliente" in _apply(api, ("i3", 50)).json()["detail"]
    _credit(api, 1000)
    assert "passa do que falta" in _apply(api, ("i1", 600)).json()["detail"]
    assert _apply(api, ("i1", 100), day=DAY(20)).status_code == 400   # antes da emissão
    assert _apply(api, ("i1", 100), day=(TODAY + timedelta(days=1)).isoformat()).status_code == 400
    assert _credit(api, 0).status_code == 400
    assert _credit(api, 5000, entry_type="DEVOLUCAO").status_code == 400  # devolução maior que o saldo
    assert _credit(api, 100, entry_type="USO").status_code == 400
    assert _balance(api) == 1100  # nada do que falhou mexeu no saldo
    assert api.get("/api/client-credits/accounts", as_user="operator").status_code == 403


def test_reversals(api):
    _seed(api)
    credit = _credit(api, 500).json()
    use = _apply(api, ("i1", 500)).json()["entries"][0]
    # crédito já usado não pode ser estornado
    assert "já foi usado" in api.post(f"/api/client-credits/entries/{credit['id']}/reverse", json={"reason": "erro"}).json()["detail"]
    assert api.post(f"/api/client-credits/entries/{use['id']}/reverse", json={"reason": " "}).status_code == 400
    # estornar a baixa reabre a fatura quitada pelo crédito
    rev = api.post(f"/api/client-credits/entries/{use['id']}/reverse", json={"reason": "fatura errada"})
    assert rev.status_code == 200 and rev.json()["entry_type"] == "ESTORNO"
    inv = _invoice(api, "i1")
    assert (inv["status"], inv["credit_applied"], inv["paid_date"]) == ("PENDENTE", 0, None)
    assert api.post(f"/api/client-credits/entries/{use['id']}/reverse", json={"reason": "de novo"}).status_code == 400
    assert api.post(f"/api/client-credits/entries/{rev.json()['id']}/reverse", json={"reason": "x"}).status_code == 400
    # agora o crédito pode ser estornado
    assert api.post(f"/api/client-credits/entries/{credit['id']}/reverse", json={"reason": "erro"}).status_code == 200
    assert _balance(api) == 0


def test_invoice_status_rules(api):
    _seed(api)
    _credit(api, 1000)
    _apply(api, ("i1", 500), ("i2", 200))
    # quitada pelo crédito: não volta a Pendente na mão
    assert api.put("/api/invoices/i1/status", params={"status": "PENDENTE"}).status_code == 400
    # parcial paga por outro meio: estornar o crédito exige reabrir antes
    assert api.put("/api/invoices/i2/status", params={"status": "PAGO"}).status_code == 200
    use_i2 = next(e for e in api.get("/api/client-credits/invoice/i2").json() if e["entry_type"] == "USO")
    assert "outro meio" in api.post(f"/api/client-credits/entries/{use_i2['id']}/reverse", json={"reason": "x"}).json()["detail"]
    # cancelar devolve o crédito
    assert api.put("/api/invoices/i1/status", params={"status": "CANCELADO"}).status_code == 200
    assert _balance(api) == 800 and _invoice(api, "i1")["credit_applied"] == 0
    # excluir também
    assert api.put("/api/invoices/i2/status", params={"status": "PENDENTE"}).status_code == 200
    assert api.delete("/api/invoices/i2").status_code == 200
    assert _balance(api) == 1000


def test_invoice_total_change_returns_excess(api):
    _seed(api)
    _credit(api, 500)
    _apply(api, ("i1", 450))
    # tira a movimentação de 300: total cai para 200, sobram 250 de crédito
    assert api.put("/api/invoices/i1", json={"movement_ids_to_remove": ["m0"]}).status_code == 200
    inv = _invoice(api, "i1")
    assert (inv["total_value"], inv["credit_applied"], inv["status"], inv["paid_by_credit"]) == (200, 200, "PAGO", True)
    assert _balance(api) == 300
    estorno = next(e for e in api.get("/api/client-credits/invoice/i1").json() if e["entry_type"] == "ESTORNO")
    assert estorno["amount"] == 250
    # devolve a movimentação: total volta a 500 e a fatura reabre com 300 em aberto
    assert api.put("/api/invoices/i1", json={"movement_ids_to_add": ["m0"]}).status_code == 200
    inv = _invoice(api, "i1")
    assert (inv["total_value"], inv["status"], inv["open_value"]) == (500, "PENDENTE", 300)
    # trocar o cliente com crédito usado é bloqueado
    assert api.put("/api/invoices/i1", json={"client_name": "OUTRO"}).status_code == 400


def test_statement_reports_and_summary(api):
    _seed(api)
    _credit(api, 1000, day=DAY(5), payment_method="TED")
    _apply(api, ("i1", 500), day=DAY(3))
    _credit(api, 100, entry_type="DEVOLUCAO", day=DAY(1))
    data = api.get("/api/client-credits/statement/cA", params={"date_from": DAY(4)}).json()
    assert data["opening_balance"] == 1000 and data["closing_balance"] == 400 and data["current_balance"] == 400
    assert [r["balance"] for r in data["rows"]] == [500, 400]
    assert (data["period_in"], data["period_out"]) == (0, 600)
    assert api.get("/api/client-credits/statement/cA/pdf").content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(api.get("/api/client-credits/statement/cA/excel").content))
    assert wb.sheetnames == ["Extrato"]
    entries = api.get("/api/client-credits/entries", params={"client_id": "cA", "entry_type": "USO"}).json()
    assert entries["total"] == 1 and entries["items"][0]["invoice_number"] == 1
    assert api.get("/api/client-credits/entries", params={"search": "1"}).json()["total"] >= 1

    # Contas a Receber e Resumo Financeiro não contam o crédito duas vezes
    rec = api.get("/api/reports/receivables").json()
    assert next(r for r in rec["rows"] if r["invoice_number"] == 1)["status"] == "PAGO"
    assert rec["summary"]["open_total"] == 900  # i2 500 + i3 400
    period = {"date_from": DAY(30), "date_to": DAY()}
    summary = api.get("/api/reports/financial-summary", params=period).json()
    ins = {r["key"]: r["total"] for r in summary["inflows"]}
    outs = {r["key"]: r["total"] for r in summary["outflows"]}
    assert ins["invoices"] == 0 and ins["client_credits"] == 1000 and outs["client_credit_refunds"] == 100
    assert next(r for r in summary["payable"] if r["key"] == "client_credit_balance")["total"] == 400


def test_service_outside_system(api):
    _seed(api)
    _credit(api, 1000)
    # baixa só de serviço, sem fatura
    r = api.post("/api/client-credits/apply", json={"client_id": "cA", "entry_date": DAY(),
                                                     "services": [{"description": "Lavagem do container", "amount": 150}]})
    assert r.status_code == 200 and r.json()["entries"][0]["entry_type"] == "SERVICO"
    assert _balance(api) == 850
    # fatura + serviço na mesma baixa; o total não pode passar do saldo
    body = {"client_id": "cA", "entry_date": DAY(), "allocations": [{"invoice_id": "i1", "amount": 500}],
            "services": [{"description": "Reparo", "amount": 400}]}
    assert "insuficiente" in api.post("/api/client-credits/apply", json=body).json()["detail"]
    body["services"][0]["amount"] = 300
    assert api.post("/api/client-credits/apply", json=body).status_code == 200
    assert _balance(api) == 50 and _invoice(api, "i1")["status"] == "PAGO"
    # serviço sem descrição ou sem valor
    for svc in ({"description": " ", "amount": 10}, {"description": "Algo", "amount": 0}):
        assert api.post("/api/client-credits/apply", json={"client_id": "cA", "entry_date": DAY(), "services": [svc]}).status_code == 400
    row = next(r for r in api.get("/api/client-credits/accounts").json() if r["client_id"] == "cA")
    assert row["total_used"] == 950
    services = api.get("/api/client-credits/entries", params={"entry_type": "SERVICO", "search": "lavagem"}).json()
    assert services["total"] == 1
    # estorno devolve ao saldo e aparece no extrato
    rev = api.post(f"/api/client-credits/entries/{services['items'][0]['id']}/reverse", json={"reason": "não foi feito"})
    assert rev.status_code == 200 and _balance(api) == 200
    st = api.get("/api/client-credits/statement/cA").json()
    assert st["closing_balance"] == 200
    assert api.get("/api/client-credits/statement/cA/pdf").content[:4] == b"%PDF"
