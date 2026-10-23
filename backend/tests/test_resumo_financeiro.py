"""Resumo Financeiro: entradas e saídas do período vindas de cada módulo,
posição do que falta receber/pagar, módulo não contratado fora do resumo, e
a data de pagamento das Diárias."""
import io

from openpyxl import load_workbook

PERIOD = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


def _insert(api, collection, *docs):
    api.run(api.db[collection].insert_many([dict(d) for d in docs]))


def _seed(api):
    _insert(api, "invoices",
            {"id": "i1", "status": "PAGO", "total_value": 1000, "paid_date": "2026-10-05"},
            {"id": "i2", "status": "PAGO", "total_value": 9999, "paid_date": "2026-09-30"},   # fora do período
            {"id": "i3", "status": "PAGO", "total_value": 50},                                # antiga, sem data
            {"id": "i4", "status": "PENDENTE", "total_value": 700})
    _insert(api, "port_service_billing_batches",
            {"id": "b1", "status": "PAGO", "net_total": 400, "total_value": 450, "paid_at": "2026-10-10T15:00:00+00:00"})
    _insert(api, "container_sales",
            {"id": "s1", "status": "RECEBIDO", "sale_value": 8000, "received_value": 7500, "received_at": "2026-10-12T15:00:00+00:00"},
            {"id": "s2", "status": "PENDENTE", "sale_value": 3000, "received_value": 1000})
    _insert(api, "checks",
            {"id": "c1", "status": "COMPENSADO", "amount": 300, "compensation_date": "2026-10-15"},
            {"id": "c2", "status": "EMITIDO", "amount": 120})
    _insert(api, "freight_payments",
            {"id": "f1", "status": "PAGO", "freight_value": 900, "paid_at": "2026-10-20T15:00:00+00:00"},
            {"id": "f2", "status": "PENDENTE", "freight_value": 250})
    _insert(api, "daily_rate_requests",
            {"id": "d1", "status": "PAGO", "total_value": 200, "updated_at": "2026-10-03T15:00:00+00:00",  # sem paid_at: usa updated_at
             "items": [{"driver_name": "A", "total": 150}, {"driver_name": "B", "total": 50}]},  # sem pagamento por item: vale o status
            {"id": "d2", "status": "PENDENTE", "total_value": 80, "items": [{"driver_name": "C", "total": 80}]})
    _insert(api, "expense_reports", {"id": "e1", "deposits": [{"amount": 500, "date": "2026-10-01", "sent_by": "Fin"}], "purchases": []})
    _insert(api, "container_purchases", {"id": "p1", "container_number": "X", "purchase_value": 5000, "entry_date": "2026-10-02T12:00:00+00:00"})


def test_inflows_outflows_and_positions(api):
    _seed(api)
    data = api.get("/api/reports/financial-summary", params=PERIOD).json()
    ins = {r["key"]: r["total"] for r in data["inflows"]}
    outs = {r["key"]: r["total"] for r in data["outflows"]}
    assert ins == {"invoices": 1000, "port_billing": 400, "container_sales": 7500, "client_credits": 0}
    assert outs == {"checks": 300, "freight": 900, "daily": 200, "expense_deposits": 500, "container_purchases": 5000, "client_credit_refunds": 0}
    s = data["summary"]
    assert s["inflow_total"] == 8900 and s["outflow_total"] == 6900 and s["balance"] == 2000
    assert {r["key"]: r["total"] for r in data["receivable"]} == {"invoices_open": 700, "port_open": 0, "sales_open": 2000}
    assert {r["key"]: r["total"] for r in data["payable"]} == {"checks_open": 120, "freight_open": 250, "daily_open": 80, "client_credit_balance": 0}
    invoices_line = next(r for r in data["inflows"] if r["key"] == "invoices")
    assert "1 fatura" in invoices_line["note"]


def test_disabled_module_is_left_out(api):
    api.run(api.db.module_config.insert_one({"disabled_modules": ["financeiro.controle_cheque", "gestao_container"]}))
    _seed(api)
    data = api.get("/api/reports/financial-summary", params=PERIOD).json()
    keys = {r["key"] for group in ("inflows", "outflows", "receivable", "payable") for r in data[group]}
    assert not keys & {"checks", "checks_open", "container_sales", "sales_open", "container_purchases"}
    assert "invoices" in keys


def test_period_required_files_and_admin_only(api):
    _seed(api)
    assert api.get("/api/reports/financial-summary").status_code == 400
    assert api.get("/api/reports/financial-summary/pdf", params=PERIOD).content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(api.get("/api/reports/financial-summary/excel", params=PERIOD).content))
    assert wb.sheetnames == ["Entradas", "Saídas", "A receber", "A pagar"]
    assert api.get("/api/reports/financial-summary", params=PERIOD, as_user="operator").status_code == 403


def test_daily_rate_paid_at(api):
    body = {"items": [{"driver_name": "JOSE", "vehicle_plate": "AAA1A11", "client_name": "C", "departure_date": "2026-10-16",
                       "daily_rate_quantity": 1, "daily_rate_value": 80}]}
    req = api.post("/api/daily-rate-requests", json=body).json()
    status = lambda s: api.put(f"/api/daily-rate-requests/{req['id']}/update-status", params={"new_status": s})
    stored = lambda: api.run(api.db.daily_rate_requests.find_one({"id": req["id"]}))
    assert status("PAGO").status_code == 200
    first = stored()["paid_at"]
    assert first
    status("PAGO")  # repetir não muda a data
    assert stored()["paid_at"] == first
    status("PENDENTE")
    assert stored()["paid_at"] is None
    assert status("PAGO").status_code == 200 and api.put("/api/daily-rate-requests/nao-existe/update-status",
                                                        params={"new_status": "PAGO"}).status_code == 404
