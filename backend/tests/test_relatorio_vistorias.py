"""Listas de Vistoria e Registro Fotográfico filtradas no banco (busca,
período em Brasília, estado) e o relatório do período das duas telas."""
import io

from openpyxl import load_workbook


def _vistoria(n, container, client, created_at, no_damage=False, items=(), photos=0, movement_id=None):
    return {"id": f"v{n}", "vistoria_number": n, "container_number": container, "client_name": client,
            "no_damage": no_damage, "damage_items": list(items), "photos": [{"url": "x"}] * photos,
            "movement_id": movement_id, "created_by": "adm", "created_by_name": "Admin", "created_at": created_at}


def _seed(api):
    api.run(api.db.container_vistorias.insert_many([
        _vistoria(1, "ABCU1111111", "CLIENTE A", "2026-10-05T15:00:00+00:00", no_damage=True, photos=3, movement_id="m1"),
        _vistoria(2, "ABCU2222222", "CLIENTE A", "2026-10-06T15:00:00+00:00", items=["Solda", "Pintura"], photos=2),
        _vistoria(3, "XYZU3333333", "CLIENTE B", "2026-10-07T15:00:00+00:00", items=["Solda"]),
        # 01/11 às 01:00 UTC ainda é 31/10 em Brasília
        _vistoria(4, "XYZU4444444", "CLIENTE B", "2026-11-01T01:00:00+00:00"),
        _vistoria(5, "OLDU5555555", "CLIENTE C", "2026-09-30T15:00:00+00:00", no_damage=True),
    ]))


PERIOD = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


def test_list_filters_run_on_the_database(api):
    _seed(api)
    get = lambda **p: api.get("/api/container-vistorias", params={"per_page": 2, **p}).json()
    data = get(search="cliente b")
    assert data["total"] == 2 and {v["vistoria_number"] for v in data["items"]} == {3, 4}
    assert get(**PERIOD)["total"] == 4
    assert get(date_from="2026-10-31", date_to="2026-10-31")["total"] == 1
    assert get(damage="COM_AVARIA")["total"] == 2
    assert get(damage="NAO_INFORMADO", **PERIOD)["total"] == 1
    assert api.get("/api/container-vistorias", params={"date_from": "31/10/2026"}).status_code == 400
    assert api.get("/api/container-inspections", params={"search": "x", **PERIOD}).json()["total"] == 0


def test_vistoria_report(api):
    _seed(api)
    data = api.get("/api/reports/container-vistorias", params=PERIOD).json()
    s = data["summary"]
    assert (s["count"], s["SEM_AVARIA"], s["COM_AVARIA"], s["NAO_INFORMADO"], s["without_photos"], s["without_eir"]) == (4, 1, 2, 1, 2, 3)
    assert data["items"][0] == {"item": "Solda", "count": 2}
    assert data["clients"][0]["client_name"] in ("CLIENTE A", "CLIENTE B")
    row4 = next(r for r in data["rows"] if r["number"] == 4)
    assert row4["day"] == "2026-10-31" and row4["pending"] == "Estado não marcado; Sem fotos; Sem EIR de entrada"
    assert api.get("/api/reports/container-vistorias/pdf", params=PERIOD, as_user="operator").content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(api.get("/api/reports/container-vistorias/excel", params=PERIOD).content))
    assert wb.sheetnames == ["Por cliente", "Avarias", "Registros"]


def test_inspection_report(api):
    api.run(api.db.container_inspections.insert_one({
        "id": "i1", "inspection_number": 1, "container_number": "ABCU1", "client_name": "C", "no_damage": False,
        "damage_items": ["Porta"], "photos": [], "created_by": "adm", "created_by_name": "Admin",
        "created_at": "2026-10-05T15:00:00+00:00"}))
    data = api.get("/api/reports/container-inspections", params=PERIOD).json()
    assert data["summary"]["COM_AVARIA"] == 1 and data["rows"][0]["pending"] == "Sem fotos"
    assert api.get("/api/reports/container-inspections/pdf", params=PERIOD).content[:4] == b"%PDF"
    assert api.get("/api/reports/container-inspections/excel", params=PERIOD).status_code == 200
