"""Conferência Ordens de Carregamento x Pagamento Frete: cada tipo de
divergência, totais por transportadora, período e arquivos."""
import io

from openpyxl import load_workbook

T = "2026-10-10T15:00:00+00:00"


def _order(id, number, status="APROVADA", route="r1", value=100.0, company="TRANSP A", created_at=T):
    return {"id": id, "order_number": number, "order_type": "COLETA", "status": status, "route_id": route,
            "route_name": "A x B" if route else None, "freight_value": value, "driver_name": "JOSE",
            "transport_company": company, "created_at": created_at}


def _payment(id, number, order_id, order_number, status="PENDENTE", value=100.0, company="TRANSP A"):
    return {"id": id, "payment_number": number, "loading_order_id": order_id, "order_number": order_number,
            "route_id": "r1", "freight_value": value, "status": status, "transport_company": company,
            "driver_name": "JOSE", "created_at": T}


def _seed(api):
    api.run(api.db.loading_orders.insert_many([
        _order("o1", 1),                                   # OK, pendente
        _order("o2", 2, value=150),                        # OK, pago
        _order("o3", 3),                                   # aprovada sem frete
        _order("o4", 4, route=None),                       # aprovada sem rota
        _order("o5", 5, value=200),                        # valor divergente
        _order("o6", 6, status="CANCELADA"),               # cancelada com frete pago
        _order("o7", 7, status="PENDENTE"),                # nada a conferir
        _order("o8", 8, created_at="2026-09-01T15:00:00+00:00"),  # fora do período
        {**_order("o9", 9), "order_type": "ENTREGA"},      # entrega: não entra
    ]))
    api.run(api.db.freight_payments.insert_many([
        _payment("f1", 1, "o1", 1),
        {**_payment("f2", 2, "o2", 2, status="PAGO", value=150), "paid_at": "2026-10-12T15:00:00+00:00"},
        _payment("f5", 5, "o5", 5, value=180),
        _payment("f6", 6, "o6", 6, status="PAGO"),
        _payment("f10", 10, "apagada", 10, company="TRANSP B"),
    ]))


PERIOD = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


def test_issues_and_totals(api):
    _seed(api)
    data = api.get("/api/reports/freight-reconciliation", params=PERIOD).json()
    issues = sorted((i["order_number"], i["issue"]) for i in data["issues"])
    assert issues == [
        (3, "Ordem aprovada sem lançamento de frete"),
        (4, "Coleta aprovada sem rota: frete não gerado"),
        (5, "Valor do frete diferente do da ordem"),
        (6, "Frete ativo com a ordem fora de Aprovada"),
        (10, "Frete ativo com a ordem excluída"),
    ]
    assert [r["order_number"] for r in data["rows"]] == [1, 2, 3, 4, 5, 6]
    assert next(r for r in data["rows"] if r["order_number"] == 1)["situation"] == "OK"
    company = data["companies"][0]
    assert company["transport_company"] == "TRANSP A"
    assert (company["total"], company["paid"], company["pending"]) == (530, 250, 280)
    assert data["summary"]["issues"] == 5


def test_company_filter_files_and_admin_only(api):
    _seed(api)
    data = api.get("/api/reports/freight-reconciliation", params={**PERIOD, "transport_company": "transp b"}).json()
    assert data["rows"] == [] and [i["order_number"] for i in data["issues"]] == [10]
    assert api.get("/api/reports/freight-reconciliation/pdf", params=PERIOD).content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(api.get("/api/reports/freight-reconciliation/excel", params=PERIOD).content))
    assert wb.sheetnames == ["Divergências", "Por transportadora", "Ordens"]
    assert api.get("/api/reports/freight-reconciliation", as_user="operator").status_code == 403


def test_period_edges(api):
    # ordem de outubro com frete lançado em novembro; frete de outubro de uma ordem de setembro
    api.run(api.db.loading_orders.insert_many([_order("a", 1), _order("b", 2, created_at="2026-09-20T15:00:00+00:00")]))
    api.run(api.db.freight_payments.insert_many([
        {**_payment("fa", 1, "a", 1), "created_at": "2026-11-03T15:00:00+00:00"},
        _payment("fb", 2, "b", 2),
    ]))
    data = api.get("/api/reports/freight-reconciliation", params=PERIOD).json()
    assert data["issues"] == []
    assert [r["order_number"] for r in data["rows"]] == [1] and data["rows"][0]["situation"] == "OK"
