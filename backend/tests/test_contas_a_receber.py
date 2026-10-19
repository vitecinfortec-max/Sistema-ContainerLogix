"""Faturas com vencimento e data de recebimento, e o relatório Contas a
Receber: faixas de atraso (do vencimento, ou da emissão quando a fatura não
tem vencimento), totais por cliente, prazo médio de recebimento, filtros e
os arquivos PDF/Excel."""
from datetime import datetime, time, timedelta

from reports import BRT_TZ
from test_api_invoices import movement, invoice

TODAY = datetime.now(BRT_TZ).date()


def _day(offset):
    return (TODAY + timedelta(days=offset)).isoformat()


def _stored(api, number, client, value, issued_offset, due_offset=None, status="PENDENTE", paid_offset=None):
    issued = datetime.combine(TODAY + timedelta(days=issued_offset), time(12), BRT_TZ)
    api.run(api.db.invoices.insert_one({
        "id": f"inv-{number}", "invoice_number": number, "client_name": client, "movement_ids": [],
        "total_value": value, "status": status, "created_at": issued.isoformat(), "created_by": "admin-1",
        "user_name": "Admin Teste", "due_date": _day(due_offset) if due_offset is not None else None,
        "paid_date": _day(paid_offset) if paid_offset is not None else None,
    }))


def test_invoice_due_date_and_paid_date(api):
    inv = api.post("/api/invoices", json={"client_name": "Cliente Um", "movement_ids": [movement(api, "ABCU1111111")["id"]],
                                          "due_date": _day(30)}).json()
    assert inv["due_date"] == _day(30) and inv["paid_date"] is None

    r = api.post("/api/invoices", json={"client_name": "X", "movement_ids": [movement(api, "ABCU2222222")["id"]], "due_date": _day(-1)})
    assert r.status_code == 400 and "vencimento" in r.json()["detail"]

    assert api.put(f"/api/invoices/{inv['id']}", json={"due_date": _day(10)}).json()["due_date"] == _day(10)
    assert api.put(f"/api/invoices/{inv['id']}", json={"due_date": ""}).json()["due_date"] is None

    assert api.put(f"/api/invoices/{inv['id']}/status", params={"status": "PAGO", "paid_date": _day(1)}).status_code == 400
    assert api.put(f"/api/invoices/{inv['id']}/status", params={"status": "PAGO"}).status_code == 200
    assert api.get(f"/api/invoices/{inv['id']}").json()["paid_date"] == _day(0)
    history = api.get(f"/api/invoices/{inv['id']}/history").json()
    assert any("paid_date" in h["changes"] for h in history)

    api.put(f"/api/invoices/{inv['id']}/status", params={"status": "PENDENTE"})
    assert api.get(f"/api/invoices/{inv['id']}").json()["paid_date"] is None


def test_aging_by_client_and_average_days(api):
    _stored(api, 1, "Alfa", 100, -10, due_offset=5)                 # a vencer
    _stored(api, 2, "Alfa", 200, -40, due_offset=-10)               # 10 dias
    _stored(api, 3, "Beta", 300, -100)                              # sem vencimento: 100 dias da emissão
    _stored(api, 4, "Beta", 400, -20, status="PAGO", paid_offset=-5)
    _stored(api, 5, "Beta", 999, -5, status="CANCELADO")            # cancelada não entra

    data = api.get("/api/reports/receivables").json()
    s = data["summary"]
    assert s["count"] == 4
    assert s["open_total"] == 600 and s["overdue_total"] == 500 and s["paid_total"] == 400
    assert s["avg_days_to_receive"] == 15
    assert s["without_due_date"] == 1
    aging = {a["key"]: a["total"] for a in s["aging"]}
    assert aging == {"current": 100, "d1_30": 200, "d31_60": 0, "d61_90": 0, "d90": 300}

    # mais atrasada primeiro, recebidas no fim
    assert [r["invoice_number"] for r in data["rows"]] == [3, 2, 1, 4]
    by_number = {r["invoice_number"]: r for r in data["rows"]}
    assert by_number[3]["aging_from_issue"] is True and by_number[3]["days_overdue"] == 100

    clients = {c["client_name"]: c for c in data["clients"]}
    assert clients["Alfa"] == {"client_name": "Alfa", "count": 2, "open": 300, "overdue": 200, "paid": 0}
    assert clients["Beta"]["paid"] == 400 and clients["Beta"]["overdue"] == 300


def test_filters(api):
    _stored(api, 1, "Alfa", 100, -10, due_offset=5)
    _stored(api, 2, "Beta", 400, -20, status="PAGO", paid_offset=-5)

    assert [r["invoice_number"] for r in api.get("/api/reports/receivables", params={"status": "PENDENTE"}).json()["rows"]] == [1]
    assert [r["invoice_number"] for r in api.get("/api/reports/receivables", params={"client_name": "bet"}).json()["rows"]] == [2]
    by_paid = api.get("/api/reports/receivables", params={"date_field": "paid", "date_from": _day(-6), "date_to": _day(0)}).json()
    assert [r["invoice_number"] for r in by_paid["rows"]] == [2]
    by_issue = api.get("/api/reports/receivables", params={"date_from": _day(-15)}).json()
    assert [r["invoice_number"] for r in by_issue["rows"]] == [1]
    assert api.get("/api/reports/receivables", params={"status": "CANCELADO"}).status_code == 422


def test_files_and_admin_only(api):
    _stored(api, 1, "Alfa", 100, -10, due_offset=5)
    _stored(api, 2, "Beta", 400, -20, status="PAGO", paid_offset=-5)
    pdf = api.get("/api/reports/receivables/pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    xlsx = api.get("/api/reports/receivables/excel")
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    from openpyxl import load_workbook
    import io
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
    assert "Alfa" in values and "Erro ao gerar Contas a Receber." not in values
    assert api.get("/api/reports/receivables/pdf", params={"client_name": "ninguém"}).content[:4] == b"%PDF"
    assert api.get("/api/reports/receivables", as_user="operator").status_code == 403
