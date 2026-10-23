"""Prestação de Contas: acerto do saldo (ressarcimento pago ao funcionário ou
devolução recebida) e o relatório do que está pendente e do que foi acertado
por período."""
import io

from openpyxl import load_workbook


def _report(api, number, start, end, purchases, deposits, status="CONCLUIDA", **extra):
    api.run(api.db.expense_reports.insert_one({
        "id": f"rep-{number}", "report_number": number, "report_number_formatted": f"2026{number:04d}",
        "period_start": start, "period_end": end, "status": status, "purchases": [], "deposits": [],
        "total_purchases": purchases, "total_deposits": deposits, "balance": round(purchases - deposits, 2),
        "created_by": "u", "created_by_name": "Ana", "created_at": "2026-10-01T12:00:00+00:00", **extra,
    }))
    return f"rep-{number}"


def _settle(api, rid, settled=True, date=None):
    body = {"settled": settled}
    if date:
        body["date"] = date
    return api.put(f"/api/expense-reports/{rid}/settlement", json=body)


def test_settle_and_undo(api):
    rid = _report(api, 1, "2026-10-01", "2026-10-15", 300, 200)
    res = _settle(api, rid, date="2026-10-20")
    assert res.status_code == 200
    assert res.json()["settled_at"] == "2026-10-20" and res.json()["settled_by_name"]

    # acertada não reabre
    reopen = api.put(f"/api/expense-reports/{rid}/status", params={"status": "EM_ANDAMENTO"})
    assert reopen.status_code == 400 and "acerto" in reopen.json()["detail"]

    assert _settle(api, rid, settled=False).json()["settled_at"] is None
    assert api.put(f"/api/expense-reports/{rid}/status", params={"status": "EM_ANDAMENTO"}).status_code == 200


def test_settle_rules(api):
    open_rid = _report(api, 1, "2026-10-01", "2026-10-15", 300, 200, status="EM_ANDAMENTO")
    assert _settle(api, open_rid).status_code == 400  # precisa estar concluída
    zero = _report(api, 2, "2026-10-01", "2026-10-15", 200, 200)
    assert _settle(api, zero).status_code == 400  # sem saldo
    rid = _report(api, 3, "2026-10-01", "2026-10-15", 100, 150)
    assert _settle(api, rid, date="20/10/2026").status_code == 400
    assert _settle(api, rid).json()["settled_at"]  # sem data: hoje
    assert _settle(api, "nao-existe").status_code == 404


def test_report_pending_and_settled_by_period(api):
    _report(api, 1, "2026-09-20", "2026-10-05", 500, 200, settled_at="2026-10-10", settled_by_name="Fin")  # ressarcido 300
    _report(api, 2, "2026-10-01", "2026-10-15", 100, 250)                                                  # devolver 150, pendente
    _report(api, 3, "2026-10-10", "2026-10-20", 400, 300, status="EM_ANDAMENTO")                          # ressarcir 100, pendente
    _report(api, 4, "2026-10-01", "2026-10-15", 200, 200)                                                  # quitada: fica de fora
    _report(api, 5, "2026-08-01", "2026-08-31", 90, 0)                                                     # fora do período

    october = {"date_from": "2026-10-01", "date_to": "2026-10-31"}
    data = api.get("/api/reports/expenses/settlements", params=october).json()
    s = data["summary"]
    assert [r["report"] for r in data["items"]] == ["20260001", "20260002", "20260003"]
    assert s["pending_reimburse"] == 100 and s["pending_return"] == 150
    assert s["settled_reimburse"] == 300 and s["settled_return"] == 0
    assert s["pending_count"] == 2 and s["settled_count"] == 1

    pending = api.get("/api/reports/expenses/settlements", params={**october, "situation": "PENDENTE"}).json()
    assert [r["report"] for r in pending["items"]] == ["20260002", "20260003"]

    # pela data do acerto só entra o que foi acertado no período
    by_settlement = api.get("/api/reports/expenses/settlements", params={**october, "period_by": "settlement"}).json()
    assert [r["report"] for r in by_settlement["items"]] == ["20260001"]
    september = {"date_from": "2026-09-01", "date_to": "2026-09-30", "period_by": "settlement"}
    assert api.get("/api/reports/expenses/settlements", params=september).json()["items"] == []
    assert api.get("/api/reports/expenses/settlements", params={"situation": "PAGO"}).status_code == 422


def test_report_files(api):
    _report(api, 1, "2026-10-01", "2026-10-15", 500, 200, settled_at="2026-10-10", settled_by_name="Fin")
    _report(api, 2, "2026-10-01", "2026-10-15", 100, 250)
    assert api.get("/api/reports/expenses/settlements/pdf").content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(api.get("/api/reports/expenses/settlements/excel").content))
    assert wb.sheetnames == ["Pendentes", "Acertadas"]
    only_pending = load_workbook(io.BytesIO(api.get("/api/reports/expenses/settlements/excel", params={"situation": "PENDENTE"}).content))
    assert only_pending.sheetnames == ["Pendentes"]
    # sem nada também gera
    empty = {"date_from": "2020-01-01", "date_to": "2020-01-31"}
    assert api.get("/api/reports/expenses/settlements/pdf", params=empty).content[:4] == b"%PDF"
    assert api.get("/api/reports/expenses/settlements", as_user="operator").status_code == 403
