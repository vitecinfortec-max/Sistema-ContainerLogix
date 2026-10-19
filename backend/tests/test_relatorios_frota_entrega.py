"""Relatório de Revisões da Frota e Relatório de Status de Entrega do período."""
import io

from openpyxl import load_workbook

PERIOD = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


def test_revision_report(api):
    api.run(api.db.vehicles.insert_many([
        {"id": "v1", "plate": "AAA1A11", "model": "SCANIA", "vehicle_type": "CAVALO", "status": "ATIVO"},
        {"id": "v2", "plate": "BBB2B22", "model": "VOLVO", "vehicle_type": "CAVALO", "status": "ATIVO"},
        {"id": "v3", "plate": "CCC3C33", "model": "CARRETA", "vehicle_type": "CARRETA", "status": "ATIVO"},
    ]))
    api.run(api.db.vehicle_revisions.insert_many([
        {"id": "r1", "revision_number": 1, "vehicle_plate": "AAA1A11", "vehicle_model": "SCANIA", "current_km": 100000,
         "next_oil_motor_km": 100500, "revision_date": "2026-10-05T12:00:00+00:00", "created_at": "2026-10-05T12:00:00+00:00",
         "oil_used": "15W40", "mechanic_name": "JOAO"},
        {"id": "r0", "revision_number": 0, "vehicle_plate": "AAA1A11", "vehicle_model": "SCANIA", "current_km": 80000,
         "next_oil_motor_km": 90000, "revision_date": "2026-08-05T12:00:00+00:00", "created_at": "2026-08-05T12:00:00+00:00",
         "oil_used": "15W40", "mechanic_name": "JOAO"},
    ]))
    data = api.get("/api/reports/vehicle-revisions", params=PERIOD).json()
    assert [r["revision_number"] for r in data["rows"]] == [1]
    assert data["vehicles"] == [{"vehicle_plate": "AAA1A11", "vehicle_model": "SCANIA", "count": 1, "last_day": "2026-10-05", "last_km": 100000}]
    assert [(m["vehicle_plate"], m["status"]) for m in data["maintenance"]] == [("AAA1A11", "DUE_SOON")]
    assert data["without_revision"] == [{"vehicle_plate": "BBB2B22", "vehicle_model": "VOLVO"}]
    assert data["summary"]["due_soon"] == 1 and data["summary"]["without_revision"] == 1
    assert api.get("/api/reports/vehicle-revisions/pdf", params=PERIOD, as_user="operator").content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(api.get("/api/reports/vehicle-revisions/excel", params=PERIOD).content))
    assert wb.sheetnames == ["Situação hoje", "Sem revisão", "Por veículo", "Revisões"]


def _item(driver, arrival="08:00", departure="10:30", start="08:30", end="10:00", done="12:00"):
    return {"driver_name": driver, "cavalo_plate": "AAA1A11", "loading_location": "X", "arrival_time": arrival,
            "departure_time": departure, "loading_start_time": start, "loading_end_time": end, "delivery_completed": done}


def test_delivery_report(api):
    api.run(api.db.delivery_statuses.insert_many([
        {"id": "d1", "status_number": 1, "status": "CONCLUIDO", "status_date": "2026-10-02", "destination_client_name": "CLI A",
         "contracting_client_name": "C", "items": [_item("JOSE"), _item("ANA", departure="09:00", done=None)]},
        {"id": "d2", "status_number": 2, "status": "ATIVO", "status_date": "2026-10-03", "destination_client_name": "CLI B",
         "contracting_client_name": "C", "items": [_item("JOSE", arrival="23:30", departure="00:30")]},
        {"id": "d3", "status_number": 3, "status": "CANCELADO", "status_date": "2026-10-04", "destination_client_name": "CLI A",
         "contracting_client_name": "C", "items": [_item("X")]},
        {"id": "d4", "status_number": 4, "status": "CONCLUIDO", "status_date": "2026-11-04", "destination_client_name": "CLI A",
         "contracting_client_name": "C", "items": [_item("X")]},
    ]))
    data = api.get("/api/reports/delivery-status", params=PERIOD).json()
    assert data["summary"]["statuses"] == 2 and data["summary"]["trips"] == 3
    assert [t["minutes_at_client"] for t in data["trips"]] == [150, 60, 60]  # 23:30 -> 00:30 passa da meia-noite
    issues = sorted((i["status_number"], i["issue"]) for i in data["issues"])
    assert issues == [(1, "Status concluído com viagem sem entrega finalizada"), (2, "Status ainda ativo de um dia que já passou")]
    jose = next(d for d in data["drivers"] if d["name"] == "JOSE")
    assert jose["trips"] == 2 and jose["avg_at_client"] == 105
    assert api.get("/api/reports/delivery-status", params={**PERIOD, "status": "CANCELADO"}).json()["summary"]["trips"] == 1
    assert api.get("/api/reports/delivery-status/pdf", params=PERIOD).content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(api.get("/api/reports/delivery-status/excel", params=PERIOD).content))
    assert wb.sheetnames == ["Pendências", "Por cliente", "Por motorista", "Viagens"]
