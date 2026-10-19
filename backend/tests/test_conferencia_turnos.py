"""Conferência de Turnos (Operacional): horários de entrada/saída dos
Serviços Portuários comparados com o turno previsto (Dia 08:30-20:00,
Noite 20:30-08:00), com as observações."""
import pytest

from routers.port_services import shift_check
from reports import fmt_shift_minutes


# ---------- Cálculo ----------

@pytest.mark.parametrize("turno, entry, exit_, entry_diff, exit_diff, balance, status", [
    # Dia no horário
    ("DIA", "08:30", "20:00", 0, 0, 0, "NO_HORARIO"),
    # chegou 15 min antes e saiu 45 min depois = 1h a mais
    ("DIA", "08:15", "20:45", 15, 45, 60, "A_MAIS"),
    # atrasou 20 min e saiu 30 min antes = 50 min a menos
    ("DIA", "08:50", "19:30", -20, -30, -50, "A_MENOS"),
    # saiu depois da meia-noite (01:00 = 5h depois das 20:00)
    ("DIA", "08:30", "01:00", 0, 300, 300, "A_MAIS"),
    # Noite no horário (saída no dia seguinte)
    ("NOITE", "20:30", "08:00", 0, 0, 0, "NO_HORARIO"),
    # chegou 20:00 (30 antes) e saiu 07:40 (20 antes) = 10 min a mais
    ("NOITE", "20:00", "07:40", 30, -20, 10, "A_MAIS"),
    # chegou depois da meia-noite (00:30 = 4h de atraso) e saiu 09:15
    ("NOITE", "00:30", "09:15", -240, 75, -165, "A_MENOS"),
    # saiu ainda antes da meia-noite (23:00 = 9h antes)
    ("NOITE", "20:30", "23:00", 0, -540, -540, "A_MENOS"),
])
def test_shift_check(turno, entry, exit_, entry_diff, exit_diff, balance, status):
    r = shift_check(turno, entry, exit_)
    assert (r["entry_diff"], r["exit_diff"], r["balance"], r["shift_status"]) == (entry_diff, exit_diff, balance, status)
    assert r["worked"] == 690 + balance  # jornada de 11h30 + saldo


def test_shift_check_missing_or_invalid_times():
    r = shift_check("DIA", "08:40", None)
    assert r["shift_status"] == "SEM_HORARIO"
    assert r["entry_diff"] == -10 and r["exit_diff"] is None and r["balance"] is None
    assert shift_check("DIA", "", "")["shift_status"] == "SEM_HORARIO"
    assert shift_check("DIA", "25:00", "20:00")["entry_diff"] is None
    assert shift_check("DIA", "08:30:00", "20:00:00")["shift_status"] == "NO_HORARIO"
    assert shift_check(None, "08:30", "20:00")["shift_status"] == "SEM_HORARIO"


def test_fmt_shift_minutes():
    assert fmt_shift_minutes(75) == "+01:15"
    assert fmt_shift_minutes(-30) == "-00:30"
    assert fmt_shift_minutes(0) == "00:00"
    assert fmt_shift_minutes(None) == "-"
    assert fmt_shift_minutes(690, signed=False) == "11:30"


# ---------- API ----------

def service(api, date, turno, entry, exit_, obs=None, client_id="c1"):
    r = api.post("/api/port-services", json={
        "client_id": client_id, "client_name": "Cliente " + client_id, "driver_name": "Motorista", "cavalo_plate": "ABC1D23",
        "service_date": date, "turno": turno, "entry_time": entry, "exit_time": exit_, "operation_value": 100,
        "observations": obs,
    })
    assert r.status_code == 200, r.text
    return r.json()


def test_report_by_period_with_observations_and_totals(api):
    service(api, "2026-10-10", "NOITE", "20:00", "08:30", obs="Ficou aguardando navio")
    service(api, "2026-10-10", "DIA", "08:30", "20:00")
    service(api, "2026-10-11", "DIA", "09:00", "20:00", obs="Atraso no portão")
    service(api, "2026-10-12", "DIA", "08:30", None)
    service(api, "2026-10-20", "DIA", "08:30", "20:00")  # fora do período

    r = api.get("/api/reports/port-service-shifts", params={"date_from": "2026-10-10", "date_to": "2026-10-12"})
    assert r.status_code == 200, r.text
    data = r.json()
    items, s = data["items"], data["summary"]
    # ordem: data, Dia antes da Noite
    assert [(i["service_date"], i["turno"]) for i in items] == [
        ("2026-10-10", "DIA"), ("2026-10-10", "NOITE"), ("2026-10-11", "DIA"), ("2026-10-12", "DIA"),
    ]
    night = items[1]
    assert night["balance"] == 60 and night["shift_status"] == "A_MAIS"
    assert night["observations"] == "Ficou aguardando navio"
    assert night["expected_entry"] == "20:30" and night["expected_exit"] == "08:00"
    assert s["count"] == 4 and s["day_count"] == 3 and s["night_count"] == 1
    assert s["on_time_count"] == 1
    assert (s["extra_count"], s["extra_minutes"]) == (1, 60)
    assert (s["less_count"], s["less_minutes"]) == (1, -30)
    assert s["missing_count"] == 1
    assert s["with_observations_count"] == 2
    assert s["balance_minutes"] == 30
    assert s["schedule"]["DIA"] == {"entry": "08:30", "exit": "20:00"}


def test_report_filters_by_turno_and_client(api):
    service(api, "2026-10-10", "DIA", "08:30", "20:00", client_id="c1")
    service(api, "2026-10-10", "NOITE", "20:30", "08:00", client_id="c2")
    params = {"date_from": "2026-10-10", "date_to": "2026-10-10"}
    assert len(api.get("/api/reports/port-service-shifts", params={**params, "turno": "NOITE"}).json()["items"]) == 1
    assert len(api.get("/api/reports/port-service-shifts", params={**params, "client_id": "c1"}).json()["items"]) == 1


def test_report_requires_valid_period(api):
    assert api.get("/api/reports/port-service-shifts").status_code == 400
    assert api.get("/api/reports/port-service-shifts", params={"date_from": "2026-10-10"}).status_code == 400
    r = api.get("/api/reports/port-service-shifts", params={"date_from": "2026-10-12", "date_to": "2026-10-10"})
    assert r.status_code == 400 and "posterior" in r.json()["detail"]
    assert api.get("/api/reports/port-service-shifts", params={"date_from": "10/10/2026", "date_to": "2026-10-12"}).status_code == 400


def test_operator_can_open_report(api):
    params = {"date_from": "2026-10-10", "date_to": "2026-10-12"}
    assert api.get("/api/reports/port-service-shifts", params=params, as_user="operator").status_code == 200


def test_pdf_and_excel(api):
    service(api, "2026-10-10", "DIA", "08:15", "20:45", obs="Observação com <caracteres> & acentuação")
    service(api, "2026-10-10", "NOITE", "20:30", None)
    params = {"date_from": "2026-10-10", "date_to": "2026-10-10"}
    pdf = api.get("/api/reports/port-service-shifts/pdf", params=params)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    xl = api.get("/api/reports/port-service-shifts/excel", params=params)
    assert xl.status_code == 200 and xl.content[:2] == b"PK"

    from io import BytesIO
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(xl.content)).active
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "+01:00" in values and "Sem horário" in values
    assert "Observação com <caracteres> & acentuação" in values
    # período vazio também gera os arquivos
    empty = {"date_from": "2026-01-01", "date_to": "2026-01-02"}
    assert api.get("/api/reports/port-service-shifts/pdf", params=empty).status_code == 200
    assert api.get("/api/reports/port-service-shifts/excel", params=empty).status_code == 200


def test_module_gate_blocks_when_disabled(api):
    api.run(api.db.module_config.insert_one({"disabled_modules": ["operacional.conferencia_turnos"]}))
    r = api.get("/api/reports/port-service-shifts", params={"date_from": "2026-10-10", "date_to": "2026-10-10"})
    assert r.status_code == 403
    # o relatório de Serviço Portuário continua liberado
    assert api.get("/api/reports/port-services/summary").status_code == 200
