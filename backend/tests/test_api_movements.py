"""
Testes de API das rotas de movimentação, Controle de Pátio, Dashboard e
relatórios, contra o app inteiro com banco em memória (fixture `api` em
conftest.py). Rodam com `pytest tests/test_api_movements.py`, sem servidor
nem Mongo no ar.

Servem de rede de segurança pra mexer na organização desses routers: fixam o
comportamento visível pela API (status, campos e regras), não a forma do
código.
"""
from datetime import timedelta, timezone

MOVEMENT = dict(
    operation_type="ENTRADA", driver_name="Motorista Um", driver_cpf="11122233344", truck_plate="AAA0A00",
    trailer_plate_1="BBB0B00", transport_company="Transportadora Um", client_name="Cliente Um",
    container_number="ABCU1234567", status="CHEIO", size_type="20DC", shipping_line="Armador Um",
    service_type="Armazenagem", service_value=150.0, currency="BRL", observations="obs de teste",
)


def create(api, **overrides):
    response = api.post("/api/movements", json={**MOVEMENT, **overrides})
    assert response.status_code == 200, response.text
    return response.json()


def seed_yard(api):
    """Dois containers em estoque e um que já saiu."""
    create(api)
    create(api, container_number="ABCU2222222", status="VAZIO", client_name="Cliente Dois")
    create(api, container_number="ABCU3333333")
    create(api, container_number="ABCU3333333", operation_type="SAIDA")


# ---------- Autenticação ----------

def test_requires_authentication(api):
    assert api.get("/api/movements", as_user=None).status_code in (401, 403)


def test_admin_only_routes_reject_common_user(api):
    assert api.get("/api/reports/billing/pdf", as_user="operator").status_code == 403


# ---------- Emissão de EIR ----------

class TestMovementCrud:
    def test_create_returns_and_stores_the_movement(self, api):
        created = create(api, container_number=" abcu1234567 ")
        assert created["container_number"] == "ABCU1234567"
        assert created["transaction_id"] == 1
        assert created["user_name"] == "Admin Teste"
        assert created["billed"] is False

        fetched = api.get(f"/api/movements/{created['id']}").json()
        for field in ("container_number", "client_name", "service_value", "currency", "observations", "driver_cpf"):
            assert fetched[field] == created[field], field

    def test_transaction_id_is_sequential(self, api):
        first = create(api)
        second = create(api, container_number="ABCU2222222")
        assert second["transaction_id"] == first["transaction_id"] + 1

    def test_second_entry_in_a_row_is_blocked(self, api):
        create(api)
        response = api.post("/api/movements", json=MOVEMENT)
        assert response.status_code == 400
        assert "duplicidade" in response.json()["detail"]

    def test_exit_after_entry_is_allowed(self, api):
        create(api)
        assert create(api, operation_type="SAIDA")["operation_type"] == "SAIDA"

    def test_invalid_payload_is_rejected(self, api):
        assert api.post("/api/movements", json={**MOVEMENT, "service_value": -5}).status_code == 422
        assert api.post("/api/movements", json={**MOVEMENT, "currency": "EUR"}).status_code == 422

    def test_update_keeps_identity_and_changes_fields(self, api):
        created = create(api)
        response = api.put(f"/api/movements/{created['id']}", json={**MOVEMENT, "seal": "LACRE9", "service_value": 200})
        assert response.status_code == 200, response.text
        updated = response.json()
        assert updated["seal"] == "LACRE9"
        assert updated["service_value"] == 200
        assert updated["transaction_id"] == created["transaction_id"]
        assert updated["created_at"] == created["created_at"]

    def test_update_unknown_movement(self, api):
        assert api.put("/api/movements/nao-existe", json=MOVEMENT).status_code == 404

    def test_delete(self, api):
        created = create(api)
        assert api.delete(f"/api/movements/{created['id']}").status_code == 200
        assert api.get(f"/api/movements/{created['id']}").status_code == 404
        assert api.delete(f"/api/movements/{created['id']}").status_code == 404

    def test_billed_movement_is_locked_for_common_user(self, api):
        created = create(api)
        api.run(api.db.movements.update_one({"id": created["id"]}, {"$set": {"billed": True}}))
        assert api.put(f"/api/movements/{created['id']}", json=MOVEMENT, as_user="operator").status_code == 403
        assert api.delete(f"/api/movements/{created['id']}", as_user="operator").status_code == 403
        assert api.put(f"/api/movements/{created['id']}", json=MOVEMENT).status_code == 200


class TestMovementQueries:
    def test_list_all(self, api):
        seed_yard(api)
        movements = api.get("/api/movements").json()
        assert len(movements) == 4
        assert {"buyer_name", "observations", "loading_order_id"} <= set(movements[0])

    def test_list_filters(self, api):
        seed_yard(api)
        assert len(api.get("/api/movements", params={"operation_type": "SAIDA"}).json()) == 1
        assert len(api.get("/api/movements", params={"status_filter": "VAZIO"}).json()) == 1
        assert len(api.get("/api/movements", params={"container_number": "abcu3333"}).json()) == 2
        assert len(api.get("/api/movements", params={"client_name": "dois"}).json()) == 1
        assert len(api.get("/api/movements", params={"transaction_id": "2"}).json()) == 1
        assert api.get("/api/movements", params={"transaction_id": "abc"}).json() == []

    def test_list_paginated(self, api):
        seed_yard(api)
        page = api.get("/api/movements", params={"page": 2, "per_page": 3}).json()
        assert (page["total"], page["pages"], page["page"], len(page["items"])) == (4, 2, 2, 1)

    def test_list_rejects_bad_pagination_and_dates(self, api):
        assert api.get("/api/movements", params={"per_page": 5000}).status_code == 422
        assert api.get("/api/movements", params={"page": 0}).status_code == 422
        assert api.get("/api/movements", params={"date_from": "12/10/2026"}).status_code == 400

    def test_current_stock_filter(self, api):
        seed_yard(api)
        in_stock = api.get("/api/movements", params={"operation_type": "ESTOQUE"}).json()
        assert sorted(m["container_number"] for m in in_stock) == ["ABCU1234567", "ABCU2222222"]

    def test_unbilled(self, api):
        seed_yard(api)
        first = api.get("/api/movements").json()[0]
        api.run(api.db.movements.update_one({"id": first["id"]}, {"$set": {"billed": True}}))
        unbilled = api.get("/api/movements/unbilled").json()
        assert len(unbilled) == 3
        assert all(m["billed"] is False for m in unbilled)

    def test_last_for_container_and_open_entry(self, api):
        seed_yard(api)
        last = api.get("/api/movements/last-for-container/ABCU3333333").json()["movement"]
        assert last["operation_type"] == "SAIDA"
        assert api.get("/api/movements/last-for-container/ZZZU0000000").json() == {"movement": None}
        assert api.get("/api/movements/open-entry/ABCU1234567").status_code == 200

    def test_find_for_invoice(self, api):
        seed_yard(api)
        assert api.get("/api/movements/find-for-invoice", params={"search": "ABCU2222222"}).status_code == 200

    def test_sync_state_changes_on_create_edit_and_delete(self, api):
        empty = api.get("/api/movements/sync-state").json()
        assert empty == {"count": 0, "last_created_at": None, "last_updated_at": None}

        created = create(api)
        after_create = api.get("/api/movements/sync-state").json()
        assert after_create["count"] == 1 and after_create["last_created_at"]

        api.put(f"/api/movements/{created['id']}", json=MOVEMENT)
        after_edit = api.get("/api/movements/sync-state").json()
        assert after_edit["last_updated_at"] and after_edit != after_create

        api.delete(f"/api/movements/{created['id']}")
        assert api.get("/api/movements/sync-state").json()["count"] == 0

    def test_voucher_pdf(self, api):
        created = create(api)
        response = api.post("/api/movements/pdf", json={"movement_ids": [created["id"]], "via": "MOTORISTA"})
        assert response.status_code == 200
        assert response.content.startswith(b"%PDF")


# ---------- Dashboard, busca e atalhos ----------

class TestDashboard:
    def test_dashboard_stats(self, api):
        seed_yard(api)
        response = api.get("/api/dashboard")
        assert response.status_code == 200, response.text
        stats = response.json()
        assert len(stats["recent_movements"]) == 4
        assert stats["recent_movements"][0]["client_name"]

    def test_dashboard_survives_legacy_quick_exit_document(self, api):
        seed_yard(api)
        api.run(api.db.movements.insert_one({
            "id": "legado-1", "transaction_id": 900, "operation_type": "SAIDA", "container_number": "ABCU1234567",
            "status": "CHEIO", "size_type": "20DC", "shipping_line": "Armador Um", "driver_name": None,
            "vehicle_plate": "CCC0C00", "transport_company_name": "Transportadora Dois",
            "created_by": "admin-1", "created_by_name": "Admin Teste",
            "created_at": "2099-01-01T00:00:00+00:00", "updated_at": None,
        }))
        response = api.get("/api/dashboard")
        assert response.status_code == 200, response.text
        assert response.json()["recent_movements"][0]["truck_plate"] == "CCC0C00"

    def test_daily_numbers_follow_the_brasilia_day(self, api):
        # Uma movimentação às 22:30 de ontem e outra à 00:01 de hoje, no
        # horário de Brasília. Em UTC as duas caem no mesmo dia (hoje).
        from reports import today_brt

        today = today_brt()
        yesterday_night = today - timedelta(minutes=90)
        base = {**MOVEMENT, "created_by": "admin-1", "user_name": "Admin Teste", "billed": False}
        api.run(api.db.movements.insert_many([
            {**base, "id": "ontem", "transaction_id": 1, "container_number": "ABCU1111111",
             "service_value": 70.0, "created_at": yesterday_night.astimezone(timezone.utc).isoformat()},
            {**base, "id": "hoje", "transaction_id": 2, "container_number": "ABCU2222222",
             "service_value": 30.0, "created_at": (today + timedelta(minutes=1)).astimezone(timezone.utc).isoformat()},
        ]))

        stats = api.get("/api/dashboard").json()
        bars = {point["date"]: point["entries"] for point in stats["daily_chart"]}
        today_key, yesterday_key = today.strftime("%Y-%m-%d"), yesterday_night.strftime("%Y-%m-%d")
        assert stats["daily_chart"][-1]["date"] == today_key
        assert bars[today_key] == 1
        assert bars[yesterday_key] == 1
        assert stats["entries_today"] == 1

        billing = {point["date"]: point["unbilled"] for point in api.get("/api/reports/billing/daily-chart").json()}
        assert billing[today_key] == 30.0
        assert billing[yesterday_key] == 70.0

    def test_other_daily_charts_end_on_the_brasilia_day(self, api):
        from reports import today_brt

        today_key = today_brt().strftime("%Y-%m-%d")
        for path in ("/api/reports/fuel-supply/daily-chart", "/api/reports/service-orders/daily-chart",
                     "/api/reports/port-services/daily-chart"):
            chart = api.get(path).json()
            assert len(chart) == 14 and chart[-1]["date"] == today_key, path

    def test_search(self, api):
        seed_yard(api)
        assert api.get("/api/search", params={"q": "ABCU2222"}).status_code == 200

    def test_alerts_summary(self, api):
        seed_yard(api)
        summary = api.get("/api/alerts/summary").json()
        assert summary["yard_over_30_days"] == 0

    def test_user_shortcuts_roundtrip(self, api):
        assert api.get("/api/user/shortcuts").status_code == 200
        assert api.put("/api/user/shortcuts", json={"shortcuts": ["movements", "billing"]}).status_code == 200
        assert "movements" in str(api.get("/api/user/shortcuts").json())


# ---------- Controle de Pátio ----------

class TestYardControl:
    def test_yard_lists_stock_and_exits(self, api):
        seed_yard(api)
        data = api.get("/api/yard-control").json()
        in_stock = sorted(c["container_number"] for c in data["containers"] if c["in_stock"])
        exited = [c["container_number"] for c in data["containers"] if not c["in_stock"]]
        assert in_stock == ["ABCU1234567", "ABCU2222222"]
        assert exited == ["ABCU3333333"]

    def test_yard_filters(self, api):
        seed_yard(api)
        only_stock = api.get("/api/yard-control", params={"movement_type": "ESTOQUE"}).json()["containers"]
        assert len(only_stock) == 2
        only_exits = api.get("/api/yard-control", params={"movement_type": "SAIDA"}).json()["containers"]
        assert [c["container_number"] for c in only_exits] == ["ABCU3333333"]
        by_client = api.get("/api/yard-control", params={"client_name": "cliente dois"}).json()["containers"]
        assert [c["container_number"] for c in by_client] == ["ABCU2222222"]
        assert api.get("/api/yard-control", params={"date_from": "ontem"}).status_code == 400

    def test_yard_exports(self, api):
        seed_yard(api)
        assert api.get("/api/yard-control/pdf").content.startswith(b"%PDF")
        assert api.get("/api/yard-control/excel").content.startswith(b"PK")

    def test_quick_exit(self, api):
        entry = create(api)
        api.run(api.db.drivers.insert_one({"id": "d1", "name": "Motorista Dois", "cpf": "99988877766"}))
        response = api.post("/api/yard-control/quick-exit", json={
            "entry_movement_id": entry["id"], "driver_id": "d1", "driver_name": "Motorista Dois",
            "vehicle_plate": "DDD0D00", "transport_company_name": "Transportadora Dois",
        })
        assert response.status_code == 200, response.text
        exit_movement = api.get("/api/movements", params={"operation_type": "SAIDA"}).json()[0]
        assert exit_movement["truck_plate"] == "DDD0D00"
        assert exit_movement["driver_cpf"] == "99988877766"
        assert exit_movement["client_name"] == "Cliente Um"
        assert exit_movement["transaction_id"] == response.json()["transaction_id"]

        again = api.post("/api/yard-control/quick-exit", json={"entry_movement_id": entry["id"]})
        assert again.status_code == 400

    def test_quick_exit_unknown_entry(self, api):
        assert api.post("/api/yard-control/quick-exit", json={"entry_movement_id": "nao-existe"}).status_code == 404

    def test_quick_exit_respects_unit_segregation(self, api):
        entry = create(api)
        api.run(api.db.unit_segregations.insert_one({
            "id": "seg-1", "status": "ATIVO", "client_name": "Outro Cliente",
            "items": [{"container_number": "ABCU1234567", "retrieved": False}],
        }))
        response = api.post("/api/yard-control/quick-exit", json={"entry_movement_id": entry["id"]})
        assert response.status_code == 400
        assert "segregado" in response.json()["detail"]


# ---------- Relatórios ----------

class TestReports:
    def test_movement_reports(self, api):
        seed_yard(api)
        assert api.get("/api/reports/pdf").content.startswith(b"%PDF")
        assert api.get("/api/reports/excel").content.startswith(b"PK")
        assert api.get("/api/reports/pdf", params={"operation_type": "ESTOQUE"}).content.startswith(b"%PDF")
        period = {"date_from": "2020-01-01", "date_to": "2099-12-31"}
        assert api.get("/api/reports/excel", params=period).content.startswith(b"PK")
        assert api.get("/api/reports/pdf", params={"date_from": "invalida"}).status_code == 400

    def test_billing_reports(self, api):
        seed_yard(api)
        assert api.get("/api/reports/billing/pdf").content.startswith(b"%PDF")
        assert api.get("/api/reports/billing/excel").content.startswith(b"PK")
        assert api.get("/api/reports/billing/daily-chart").status_code == 200
        ids = [m["id"] for m in api.get("/api/movements").json()]
        assert api.post("/api/billing/report", json={"movement_ids": ids}).content.startswith(b"PK")

    def test_storage_overage_reports(self, api):
        seed_yard(api)
        assert api.get("/api/reports/storage-overage/pdf").content.startswith(b"%PDF")
        assert api.get("/api/reports/storage-overage/excel").content.startswith(b"PK")

    def test_storage_overage_counts_from_the_brasilia_entry_day(self, api):
        # Entrada em 01/10 às 22:30 de Brasília (já é 02/10 em UTC), free time de 5 dias.
        api.run(api.db.movements.insert_one({
            **MOVEMENT, "id": "diaria-1", "transaction_id": 500,
            "created_at": "2026-10-02T01:30:00+00:00", "created_by": "admin-1", "user_name": "Admin Teste",
        }))
        api.run(api.db.service_price_entries.insert_one({
            "id": "tab-1", "billing_type": "DIARIA", "status": "ATIVO", "client_name": "Cliente Um",
            "service_type_name": "Diária de Armazenagem", "free_time_days": 5, "value": 10.0, "currency": "BRL",
        }))
        from datetime import date

        from routers.movement_reports import _compute_storage_overage_charges
        charges = api.run(_compute_storage_overage_charges(None, None, None, date(2026, 10, 12)))
        assert [c["container_number"] for c in charges] == ["ABCU1234567"]
        # 01/10 + 5 dias de free time = 06/10; de 06/10 a 12/10 são 6 diárias.
        assert charges[0]["extra_days"] == 6
        assert charges[0]["service_value"] == 60.0


# ---------- Lista paginada no servidor (tela Gate) ----------

class TestServerSidePagination:
    def seed(self, api, count=23):
        for i in range(count):
            create(
                api, container_number=f"ABCU{i:07d}", client_name="Cliente Par" if i % 2 == 0 else "cliente ímpar",
                status="CHEIO" if i % 3 == 0 else "VAZIO", driver_name=f"Motorista {i:02d}",
            )
        create(api, container_number="ABCU0000000", operation_type="SAIDA", client_name="Cliente Par")

    def page(self, api, **params):
        response = api.get("/api/movements", params={"per_page": 10, **params})
        assert response.status_code == 200, response.text
        return response.json()

    def test_stats_cover_the_whole_filter_not_the_page(self, api):
        self.seed(api)
        data = self.page(api)
        assert len(data["items"]) == 10
        assert data["stats"] == {"total": 24, "entries": 23, "exits": 1, "full": 9, "empty": 15}

        filtered = self.page(api, client_name="cliente par")
        assert filtered["total"] == filtered["stats"]["total"] == 13
        assert filtered["stats"]["entries"] == 12 and filtered["stats"]["exits"] == 1

    def test_pages_do_not_repeat_or_skip_records(self, api):
        self.seed(api)
        seen = []
        for number in (1, 2, 3):
            seen += [m["id"] for m in self.page(api, page=number, sort_by="status", sort_dir="asc")["items"]]
        assert len(seen) == 24 and len(set(seen)) == 24

    def test_sorting(self, api):
        self.seed(api)
        newest_first = self.page(api, sort_by="transaction_id", sort_dir="desc")["items"]
        assert [m["transaction_id"] for m in newest_first[:3]] == [24, 23, 22]
        oldest_first = self.page(api, sort_by="transaction_id", sort_dir="asc")["items"]
        assert [m["transaction_id"] for m in oldest_first[:3]] == [1, 2, 3]
        by_driver = self.page(api, sort_by="driver_name", sort_dir="asc")["items"]
        assert by_driver[0]["driver_name"] == "Motorista 00"

    def test_unknown_sort_column_falls_back_to_newest_first(self, api):
        self.seed(api)
        items = self.page(api, sort_by="senha", sort_dir="asc")["items"]
        assert items[0]["transaction_id"] == 24
        assert api.get("/api/movements", params={"per_page": 10, "sort_dir": "torto"}).status_code == 422

    def test_record_number_search_matches_by_prefix(self, api):
        self.seed(api)
        assert sorted(m["transaction_id"] for m in self.page(api, search_id="2")["items"]) == [2, 20, 21, 22, 23, 24]
        assert [m["transaction_id"] for m in self.page(api, search_id="#17")["items"]] == [17]
        assert self.page(api, search_id="abc")["total"] == 0
        assert self.page(api, search_id="  ")["total"] == 24

    def test_exact_transaction_id_filter_is_unchanged(self, api):
        # Usado pela tela de Faturamento ("Editar Fatura")
        self.seed(api)
        assert [m["transaction_id"] for m in api.get("/api/movements", params={"transaction_id": "2"}).json()] == [2]

    def test_combined_filters(self, api):
        self.seed(api)
        data = self.page(api, operation_type="ENTRADA", status_filter="CHEIO", client_name="cliente par", driver_name="motorista 1")
        assert data["total"] == len(data["items"]) == 2
        assert all(m["status"] == "CHEIO" and m["operation_type"] == "ENTRADA" for m in data["items"])

    def test_page_beyond_the_end_is_empty_but_reports_totals(self, api):
        self.seed(api)
        data = self.page(api, page=9)
        assert data["items"] == [] and data["total"] == 24 and data["pages"] == 3


# ---------- Valor do container e data de fabricação ----------

class TestContainerValueAndManufactureDate:
    ENTRY = {"container_value": 12500.5, "container_value_currency": "USD", "manufacture_date": "2019-05"}

    def test_entry_stores_both_fields(self, api):
        created = create(api, **self.ENTRY)
        fetched = api.get(f"/api/movements/{created['id']}").json()
        for field, value in self.ENTRY.items():
            assert created[field] == fetched[field] == value, field

    def test_both_fields_are_optional(self, api):
        created = create(api)
        assert created["container_value"] is None
        assert created["manufacture_date"] is None
        assert created["container_value_currency"] == "BRL"

    def test_empty_status_entry_also_stores_them(self, api):
        created = create(api, status="VAZIO", **self.ENTRY)
        assert created["container_value"] == 12500.5 and created["manufacture_date"] == "2019-05"

    def test_exit_pulls_the_values_saved_on_the_entry(self, api):
        create(api, **self.ENTRY)
        exit_movement = create(api, operation_type="SAIDA")
        for field, value in self.ENTRY.items():
            assert exit_movement[field] == value, field

    def test_values_pulled_from_the_entry_are_locked_on_the_exit(self, api):
        create(api, **self.ENTRY)
        exit_movement = create(
            api, operation_type="SAIDA", container_value=1, container_value_currency="BRL", manufacture_date="2001-01",
        )
        for field, value in self.ENTRY.items():
            assert exit_movement[field] == value, field

        # Editar a saída também não consegue trocar o que veio da entrada.
        edited = api.put(f"/api/movements/{exit_movement['id']}", json={
            **MOVEMENT, "operation_type": "SAIDA", "container_value": 2, "manufacture_date": "2002-02",
        }).json()
        for field, value in self.ENTRY.items():
            assert edited[field] == value, field

    def test_what_the_entry_left_blank_can_be_filled_on_the_exit(self, api):
        create(api, manufacture_date="2019-05")
        exit_movement = create(api, operation_type="SAIDA", container_value=900, container_value_currency="USD")
        assert exit_movement["manufacture_date"] == "2019-05"
        assert exit_movement["container_value"] == 900 and exit_movement["container_value_currency"] == "USD"

    def test_exit_without_entry_keeps_what_was_typed(self, api):
        exit_movement = create(api, operation_type="SAIDA", container_value=700, manufacture_date="2015-11")
        assert exit_movement["container_value"] == 700 and exit_movement["manufacture_date"] == "2015-11"

    def test_exit_uses_the_latest_entry_of_the_container(self, api):
        create(api, container_value=100, manufacture_date="2010-01")
        create(api, operation_type="SAIDA")
        create(api, container_value=250, manufacture_date="2010-01")
        assert create(api, operation_type="SAIDA")["container_value"] == 250

    def test_entry_can_be_corrected_later(self, api):
        entry = create(api, **self.ENTRY)
        edited = api.put(f"/api/movements/{entry['id']}", json={**MOVEMENT, "container_value": 9000, "manufacture_date": "2020-12"})
        assert edited.status_code == 200, edited.text
        assert edited.json()["container_value"] == 9000 and edited.json()["manufacture_date"] == "2020-12"

    def test_quick_exit_inherits_them(self, api):
        entry = create(api, **self.ENTRY)
        response = api.post("/api/yard-control/quick-exit", json={"entry_movement_id": entry["id"]})
        assert response.status_code == 200, response.text
        exit_movement = api.get("/api/movements", params={"operation_type": "SAIDA"}).json()[0]
        assert exit_movement["container_value"] == 12500.5 and exit_movement["manufacture_date"] == "2019-05"

    def test_open_entry_lookup_returns_them_for_the_form(self, api):
        create(api, **self.ENTRY)
        entry = api.get("/api/movements/open-entry/ABCU1234567").json()["entry"]
        assert entry["container_value"] == 12500.5 and entry["manufacture_date"] == "2019-05"

    def test_locked_fields_lookup_returns_only_what_the_entry_filled(self, api):
        assert api.get("/api/movements/entry-locked-fields/ABCU1234567").json() == {}
        create(api, manufacture_date="2019-05")
        assert api.get("/api/movements/entry-locked-fields/abcu1234567").json() == {"manufacture_date": "2019-05"}

    def test_locked_fields_lookup_for_an_exit_being_edited_ignores_later_entries(self, api):
        create(api, **self.ENTRY)
        exit_movement = create(api, operation_type="SAIDA")
        create(api, container_value=1, manufacture_date="2001-01")
        locked = api.get(
            "/api/movements/entry-locked-fields/ABCU1234567", params={"exit_id": exit_movement["id"]}
        ).json()
        assert locked == self.ENTRY

    def test_movements_report_shows_both_columns(self, api):
        import io
        from openpyxl import load_workbook
        from pypdf import PdfReader

        create(api, **self.ENTRY)
        create(api, container_number="ABCU2222222", container_value=8000, manufacture_date="2021-11")
        create(api, container_number="ABCU3333333")

        excel = api.get("/api/reports/excel")
        assert excel.status_code == 200, excel.text
        rows = [[c for c in row] for row in load_workbook(io.BytesIO(excel.content)).active.iter_rows(values_only=True)]
        header = next(r for r in rows if "Valor Container" in r)
        value_col, date_col, container_col = (header.index(h) for h in ("Valor Container", "Fabricação", "Nº Container"))
        by_container = {r[container_col]: (r[value_col], r[date_col]) for r in rows if r is not header}
        assert by_container["ABCU1234567"] == ("$ 12,500.50", "05/2019")
        assert by_container["ABCU2222222"] == ("R$ 8.000,00", "11/2021")
        assert by_container["ABCU3333333"] == ("-", "-")

        pdf = api.get("/api/reports/pdf")
        assert pdf.status_code == 200, pdf.text
        reader = PdfReader(io.BytesIO(pdf.content))
        assert len(reader.pages) == 1
        # O cabeçalho da tabela sai em maiúsculas e pode quebrar em duas linhas
        text = " ".join(reader.pages[0].extract_text().upper().split())
        for expected in ("VALOR CONTAINER", "FABRICAÇÃO", "$ 12,500.50", "05/2019", "R$ 8.000,00", "11/2021"):
            assert expected in text, expected

    def test_daily_email_report_leaves_both_columns_out(self):
        import io
        from pypdf import PdfReader
        from reports import generate_pdf_report

        movement = {**MOVEMENT, **self.ENTRY, "transaction_id": 1, "created_at": "2026-10-12T12:00:00+00:00"}
        text = PdfReader(io.BytesIO(generate_pdf_report([movement], show_container_value=False))).pages[0].extract_text()
        assert "Valor Container" not in text and "12,500.50" not in text

    def test_validation(self, api):
        assert api.post("/api/movements", json={**MOVEMENT, "container_value": -1}).status_code == 422
        assert api.post("/api/movements", json={**MOVEMENT, "container_value_currency": "EUR"}).status_code == 422
        assert api.post("/api/movements", json={**MOVEMENT, "manufacture_date": "05/2019"}).status_code == 422
        assert api.post("/api/movements", json={**MOVEMENT, "manufacture_date": "2019-13"}).status_code == 422
        # Data completa é aceita e guardada como mês/ano; texto vazio vira "não informado".
        assert create(api, manufacture_date="2019-05-17")["manufacture_date"] == "2019-05"
        assert create(api, container_number="ABCU2222222", manufacture_date="")["manufacture_date"] is None
