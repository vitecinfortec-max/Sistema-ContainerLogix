"""Relatórios de conferência do Financeiro: Diárias (por motorista e item a
item, pela data de saída) e Prestação de Contas (compras por fornecedor,
depósitos, saldo e compras sem recibo), com filtros, PDF e Excel."""
import io

from openpyxl import load_workbook


def _daily(api, *items, status=None):
    body = {"items": [{"driver_name": d, "vehicle_plate": "AAA1A11", "client_name": c, "departure_date": day,
                       "daily_rate_quantity": qty, "daily_rate_value": 100, "lunch_value": 20} for d, c, day, qty in items]}
    req = api.post("/api/daily-rate-requests", json=body).json()
    if status:
        api.put(f"/api/daily-rate-requests/{req['id']}/update-status", params={"new_status": status})
    return req


def _expense(api, number, status="EM_ANDAMENTO", purchases=(), deposits=()):
    api.run(api.db.expense_reports.insert_one({
        "id": f"rep-{number}", "report_number": number, "report_number_formatted": f"2026{number:04d}", "status": status,
        "period_start": "2026-10-01", "period_end": "2026-10-31",
        "purchases": [{"item_id": f"p{number}-{i}", "supplier_name": s, "purchase_date": day, "amount": amount,
                       "observation": obs, "receipts": [{"id": "r", "url": "/x.png"}] if receipt else []}
                      for i, (s, day, amount, obs, receipt) in enumerate(purchases)],
        "deposits": [{"amount": a, "date": day, "sent_by": who} for a, day, who in deposits],
    }))


def _sheet_values(content):
    wb = load_workbook(io.BytesIO(content))
    return {ws.title: [c for row in ws.iter_rows(values_only=True) for c in row if c is not None] for ws in wb.worksheets}


def test_daily_rates_by_driver_and_period(api):
    _daily(api, ("JOSE", "Alfa", "2026-10-05", 2), ("MARIA", "Beta", "2026-10-06", 1))
    _daily(api, ("JOSE", "Beta", "2026-10-20", 1), status="PAGO")
    _daily(api, ("JOSE", "Alfa", "2026-10-07", 5), status="CANCELADO")  # cancelada fica de fora

    data = api.get("/api/reports/daily-rates").json()
    s = data["summary"]
    assert s["count"] == 3 and s["total"] == 460 and s["paid"] == 120 and s["pending"] == 340
    assert s["daily"] == 400 and s["lunch"] == 60
    jose = data["drivers"][0]
    assert jose["driver_name"] == "JOSE" and jose["trips"] == 2 and jose["total"] == 340 and jose["pending"] == 220

    in_period = api.get("/api/reports/daily-rates", params={"date_from": "2026-10-01", "date_to": "2026-10-10"}).json()
    assert [i["driver_name"] for i in in_period["items"]] == ["JOSE", "MARIA"]
    assert api.get("/api/reports/daily-rates", params={"client": "bet"}).json()["summary"]["count"] == 2
    assert api.get("/api/reports/daily-rates", params={"status": "CANCELADO"}).json()["summary"]["total"] == 520


def test_daily_rates_files(api):
    _daily(api, ("JOSE", "Alfa", "2026-10-05", 2))
    assert api.get("/api/reports/daily-rates/pdf").content[:4] == b"%PDF"
    sheets = _sheet_values(api.get("/api/reports/daily-rates/excel").content)
    assert list(sheets) == ["Por motorista", "Itens"]
    assert "JOSE" in sheets["Por motorista"] and 220 in sheets["Itens"]
    # sem nenhum dado também gera os arquivos
    assert api.get("/api/reports/daily-rates/pdf", params={"driver": "ninguém"}).content[:4] == b"%PDF"
    assert "Erro ao gerar" not in str(_sheet_values(api.get("/api/reports/daily-rates/excel", params={"driver": "x"}).content))


def test_expenses_suppliers_and_missing_receipts(api):
    _expense(api, 1, purchases=[("Posto A", "2026-10-02", 100, "Diesel", True), ("Papelaria", "2026-10-03", 30, "Canetas", False)],
             deposits=[(200, "2026-10-01", "Financeiro")])
    _expense(api, 2, status="CONCLUIDA", purchases=[("Posto A", "2026-11-02", 50, None, False)])

    data = api.get("/api/reports/expenses").json()
    s = data["summary"]
    assert s["total_purchases"] == 180 and s["total_deposits"] == 200 and s["balance"] == -20
    assert s["without_receipt_count"] == 2 and s["without_receipt_total"] == 80
    posto = data["suppliers"][0]
    assert posto == {"supplier_name": "Posto A", "count": 2, "total": 150, "without_receipt": 1, "without_receipt_total": 50}

    october = api.get("/api/reports/expenses", params={"date_from": "2026-10-01", "date_to": "2026-10-31"}).json()
    assert october["summary"]["total_purchases"] == 130
    missing = api.get("/api/reports/expenses", params={"only_without_receipt": "true"}).json()
    assert [p["supplier_name"] for p in missing["purchases"]] == ["Papelaria", "Posto A"] and missing["deposits"] == []
    assert api.get("/api/reports/expenses", params={"status": "CONCLUIDA"}).json()["summary"]["total_purchases"] == 50


def test_expenses_files_and_admin_only(api):
    _expense(api, 1, purchases=[("Posto A", "2026-10-02", 100, "Diesel", False)], deposits=[(200, "2026-10-01", "Financeiro")])
    assert api.get("/api/reports/expenses/pdf").content[:4] == b"%PDF"
    sheets = _sheet_values(api.get("/api/reports/expenses/excel").content)
    assert list(sheets) == ["Por grupo", "Maiores compras", "Sem recibo", "Por fornecedor", "Compras", "Depósitos"]
    assert "Posto A" in sheets["Sem recibo"]
    for path in ("/api/reports/expenses", "/api/reports/daily-rates"):
        assert api.get(path, as_user="operator").status_code == 403
