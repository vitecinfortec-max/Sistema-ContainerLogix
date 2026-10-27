"""Relatório de Abastecimento (27/10/2026): sai o Nº da Ordem de
Abastecimento de cada lançamento, pra conferir com as ordens emitidas."""
import io

from openpyxl import load_workbook


def _vehicle(api, plate="ABC1D23"):
    api.run(api.db.vehicles.insert_one({
        "id": f"v-{plate}", "plate": plate, "model": "FH 540", "vehicle_type": "CAVALO", "status": "ATIVO",
        "created_by": "admin-1", "created_by_name": "Admin", "created_at": "2026-10-01T10:00:00+00:00",
    }))
    return f"v-{plate}"


def _supply(api, vid, date, reading, **extra):
    body = {
        "equipment_id": vid, "equipment_plate": "ABC1D23", "supply_date": date, "entry_date": date,
        "reading": reading, "source": "POSTO_EXTERNO", "supplier_name": "Posto X", "fuel_type": "DIESEL_S10",
        "liters": 100, "unit_price": 6, "gross_value": 600, "full_tank": True, **extra,
    }
    r = api.post("/api/fuel-supplies", json=body)
    assert r.status_code == 200, r.text
    return r.json()


def _setup(api):
    vid = _vehicle(api)
    order = api.post("/api/fuel-supply-orders", json={
        "order_date": "2026-10-01", "equipment_plate": "ABC1D23", "fuel_type": "DIESEL_S10", "supply_mode": "LITROS", "liters": 100,
    }).json()
    linked = _supply(api, vid, "2026-10-01", 1000, fuel_supply_order_id=order["id"], supply_order=f"Nº {order['order_number']}")
    typed = _supply(api, vid, "2026-10-02", 1500, supply_order="Nº 4521")  # lançamento antigo / digitado à mão
    without = _supply(api, vid, "2026-10-03", 2000)
    return order, linked, typed, without


def test_summary_rows_bring_order_number(api):
    order, linked, typed, without = _setup(api)
    rows = {r["id"]: r for r in api.get("/api/reports/fuel-supply/summary").json()["rows"]}
    assert rows[linked["id"]]["order_number_label"] == str(order["order_number"])
    assert rows[typed["id"]]["order_number_label"] == "4521"
    assert rows[without["id"]]["order_number_label"] is None


def test_linked_order_number_wins_over_text(api):
    order, linked, _, _ = _setup(api)
    api.put(f"/api/fuel-supplies/{linked['id']}", json={**linked, "supply_order": "texto qualquer"})
    rows = {r["id"]: r for r in api.get("/api/reports/fuel-supply/summary").json()["rows"]}
    assert rows[linked["id"]]["order_number_label"] == str(order["order_number"])


def test_excel_and_pdf_have_order_column(api):
    order, _, _, _ = _setup(api)
    r = api.get("/api/reports/fuel-supply/excel")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    values = [[c for c in row if c is not None] for row in ws.iter_rows(values_only=True)]
    header = next(row for row in values if "Equipamento" in row)
    assert header[0] == "Ordem"
    column = [row[0] for row in values[values.index(header) + 1:] if row]
    assert order["order_number"] in column and "-" in column and 4521 in column

    pdf = api.get("/api/reports/fuel-supply/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
