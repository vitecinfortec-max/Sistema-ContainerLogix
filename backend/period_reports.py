"""Montagem comum dos relatórios de conferência de um período (PDF e Excel).

Cada relatório monta os dados uma vez e descreve as tabelas como Section;
daqui saem o PDF (cabeçalho padrão, quadro de indicadores, uma tabela por
seção) e o Excel (primeira seção com o layout padrão das planilhas, as
demais em abas próprias), com a mesma identidade visual dos outros
documentos do sistema (reports.py).
"""
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import List, Optional
import io
import logging

from fastapi.responses import StreamingResponse

from reports import BRT_TZ, merge_company, format_currency
from shared import get_company_settings, load_logo_buffer

logger = logging.getLogger(__name__)

MONEY = "money"
DAY = "day"
INT = "int"
TEXT = "text"


@dataclass
class Column:
    label: str
    kind: str = TEXT  # money | day | int | text
    width: float = 1.0  # peso relativo da largura
    align: Optional[str] = None  # padrão: dinheiro à direita, data/número ao centro

    @property
    def pdf_align(self):
        return self.align or {MONEY: "right", DAY: "center", INT: "center"}.get(self.kind, "left")


@dataclass
class Section:
    title: str
    columns: List[Column]
    rows: list
    total: Optional[list] = None  # linha de total (mesmo número de colunas; None = vazio)
    empty_text: str = "Nada a mostrar."
    sheet_name: Optional[str] = None


def fmt_day(value) -> str:
    if not value:
        return "-"
    if isinstance(value, str):
        value = date.fromisoformat(value[:10])
    return value.strftime("%d/%m/%Y")


def period_text(label: str, date_from: Optional[str], date_to: Optional[str]) -> Optional[str]:
    if not date_from and not date_to:
        return None
    return f"{label}: {fmt_day(date_from) if date_from else 'início'} a {fmt_day(date_to) if date_to else 'hoje'}"


def _pdf_text(value, kind):
    if value is None or value == "":
        return "-"
    if kind == DAY:
        return fmt_day(value)
    if isinstance(value, str):  # ex.: "TOTAL" na primeira coluna da linha de total
        return value
    if kind == MONEY:
        return format_currency(value, "BRL")
    if kind == INT:
        if float(value).is_integer():
            return f"{int(value):,}".replace(",", ".")
        return f"{value:.2f}".rstrip("0").replace(".", ",")
    return str(value)


async def pdf_response(title: str, subtitle: Optional[str], info: list, sections: List[Section], filename: str, user_name: str,
                       dashboard=None):
    """`dashboard(width)`: flowables extras (gráficos) entre os indicadores e as tabelas."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reports import (
        _build_pdf_header, _make_pdf_footer, _pdf_info_grid, _pdf_section_title, _pdf_header_cells,
        _pdf_cell_factory, _pdf_table_style, _pdf_empty_state,
    )

    company = merge_company(await get_company_settings())
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=12*mm, leftMargin=12*mm, topMargin=10*mm, bottomMargin=14*mm)
    width = doc.width
    styles = getSampleStyleSheet()
    elements = list(_build_pdf_header(styles, load_logo_buffer(company), title, generation_info=subtitle, company=company, content_width=width))
    info = list(info) + [("Gerado em", datetime.now(BRT_TZ).strftime("%d/%m/%Y %H:%M")), ("Gerado por", user_name)]
    elements.append(_pdf_info_grid(info, width, cols=4))
    if dashboard:
        elements.extend(dashboard(width))

    cell = _pdf_cell_factory(styles, font_size=7.5)
    for section in sections:
        elements.extend(_pdf_section_title(section.title, width))
        if not section.rows:
            elements.append(_pdf_empty_state(section.empty_text, width))
            continue
        data = [_pdf_header_cells([c.label for c in section.columns], font_size=7.5)]
        for row in section.rows:
            data.append([cell(_pdf_text(v, c.kind), c.pdf_align) for v, c in zip(row, section.columns)])
        if section.total:
            data.append([cell(_pdf_text(v, c.kind), c.pdf_align, bold=True) if v is not None else cell("", markup="")
                         for v, c in zip(section.total, section.columns)])
        scale = width / sum(c.width for c in section.columns)
        table = Table(data, colWidths=[c.width * scale for c in section.columns], repeatRows=1)
        table.setStyle(TableStyle(_pdf_table_style(total_row=bool(section.total))))
        elements.append(table)
        elements.append(Spacer(1, 4))

    footer = _make_pdf_footer(company["name"])
    doc.build(elements, onFirstPage=footer, onLaterPages=footer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}.pdf"})


def _xl_value(value, kind):
    from reports import _xl_date
    if value in (None, ""):
        return None if kind in (MONEY, INT) else "-"
    if kind == DAY:
        return _xl_date(value)
    return value


async def excel_response(title: str, stats: str, sections: List[Section], filename: str):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from reports import _bsoft_style_excel, _xl_extend_print_area, _xl_error_workbook, XL_DATE, download_logo

    formats = {MONEY: "R$ #,##0.00", DAY: XL_DATE, INT: "0"}
    try:
        company = await get_company_settings()
        company_name = merge_company(company)["name"]
        wb = Workbook()
        for index, section in enumerate(sections):
            ws = wb.active if index == 0 else wb.create_sheet()
            ws.title = (section.sheet_name or section.title)[:31]
            rows = [[_xl_value(v, c.kind) for v, c in zip(row, section.columns)] for row in section.rows]
            total_cols = None
            if section.total:
                total_cols = [i for i, (v, c) in enumerate(zip(section.total, section.columns)) if c.kind == MONEY and v is not None] or None
            widths = {get_column_letter(i + 2): max(10, min(48, round(c.width * 5))) for i, c in enumerate(section.columns)}
            _bsoft_style_excel(
                ws, f"{title} - {section.title}" if index else title, stats, [c.label for c in section.columns], rows, widths,
                center_cols={i for i, c in enumerate(section.columns) if c.pdf_align == "center"},
                number_fmt_cols={i: formats[c.kind] for i, c in enumerate(section.columns) if c.kind in formats},
                total_col=total_cols, stats_text=stats, company_name=company_name, logo_buffer=download_logo(company),
                company=company, autofilter=bool(rows), empty_message=section.empty_text,
            )
            last_row = 9 + max(len(rows), 1) + (1 if total_cols else 0)
            _xl_extend_print_area(ws, get_column_letter(len(section.columns) + 1), last_row)
        buffer = io.BytesIO()
        wb.save(buffer)
        content = buffer.getvalue()
    except Exception:
        logger.exception("Erro ao gerar o Excel: %s", title)
        content = _xl_error_workbook(f"Erro ao gerar {title}.")
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"},
    )
